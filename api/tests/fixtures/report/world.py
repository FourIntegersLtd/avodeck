"""A case for slice F's tests: the named expert and other case roles, an outsider, three numbered pages
(the expert has read pages 1 and 2, not page 3), one confirmed question, a verified fact on each page,
the expert's Word template (sealed as slice C stores it) and one AI call for the methodology appendix.

Slice B's read log stands in as `Reads` (as in slice E's tests). Fact values and opinion words are short
made-up words; nothing here is a real record. The case runs under whichever pack combination the test names.
"""

from __future__ import annotations

import base64
import io
import uuid
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

import docx
import psycopg
from fastapi import FastAPI
from psycopg import AsyncConnection
from psycopg.types.json import Jsonb

from app.core.contracts import opinions
from app.core.crypto import CaseCrypto, LocalKeyProvider, SecretBox
from app.core.db import Database
from app.core.packs import default_library
from app.core.storage import CaseFileStorage
from tests.support.contracts import add_fact, add_question
from tests.support.packs import Combo
from tests.support.seed import Seed

TEMPLATE = Path(__file__).resolve().parents[1] / "setup" / "template.docx"
PLAN = "review 3/52 if no better"
RESULT = "Ferritin 9 ug/L"
LATER = "seen again in clinic"
MODEL = "zz.model-in-region-v1"
CRYPTO = CaseCrypto(LocalKeyProvider.generate())


@dataclass
class Reads:
    """In-memory read log: `(case_id, page_id, page_no, person_id)` rows standing in for slice B."""

    rows: set[tuple[uuid.UUID, uuid.UUID, int, uuid.UUID]] = field(default_factory=set)

    async def has_read(self, conn: AsyncConnection, case_id: uuid.UUID, page_no: int, person_id: uuid.UUID) -> bool:
        return any(r[0] == case_id and r[2] == page_no and r[3] == person_id for r in self.rows)

    async def read_page_ids(self, conn: AsyncConnection, case_id: uuid.UUID, person_id: uuid.UUID) -> set[uuid.UUID]:
        return {r[1] for r in self.rows if r[0] == case_id and r[3] == person_id}


@dataclass
class World:
    combo: Combo
    org: uuid.UUID
    case: uuid.UUID
    expert: uuid.UUID
    co_expert: uuid.UUID
    assistant: uuid.UUID
    viewer: uuid.UUID
    outsider: uuid.UUID
    document: uuid.UUID
    pages: list[uuid.UUID]
    question: uuid.UUID
    plan: uuid.UUID
    result: uuid.UUID
    unread: uuid.UUID

    def url(self, tail: str) -> str:
        return f"/v1/cases/{self.case}/{tail}"


def build(seed: Seed, admin: psycopg.Connection, combo: Combo, reads: Reads) -> World:
    expert, co_expert, assistant, viewer, outsider = (seed.person() for _ in range(5))
    org = seed.org(
        kind="team", members={expert: "expert", co_expert: "expert", assistant: "assistant", viewer: "expert"}
    )
    seed.org(kind="personal", members={outsider: "owner"})
    case = seed.case(
        org,
        expert,
        status="active",
        members={co_expert: "co_expert", assistant: "assistant", viewer: "viewer"},
        packs=combo.case_row(),
    )
    document = seed.document(case)
    pages = [seed.page(case, document, n) for n in (1, 2, 3)]
    for n, page in enumerate(pages[:2], start=1):
        reads.rows.add((case, page, n, expert))
    admin.execute(
        "insert into app.ai_calls (case_id, task_id, model_id, region, prompt_version, input_tokens, output_tokens,"
        " outcome) values (%s, 'T-07b', %s, 'zz', '2026-10-16.1', 1, 1, 'ok')",
        (case, MODEL),
    )
    return World(
        combo=combo,
        org=org,
        case=case,
        expert=expert,
        co_expert=co_expert,
        assistant=assistant,
        viewer=viewer,
        outsider=outsider,
        document=document,
        pages=pages,
        question=add_question(admin, case, status="confirmed"),
        plan=add_fact(admin, case, pages[0], status="verified", value=PLAN),
        result=add_fact(admin, case, pages[1], status="verified", value=RESULT),
        unread=add_fact(admin, case, pages[2], status="verified", value=LATER),
    )


def cite(fact: uuid.UUID, page: uuid.UUID, quote: str) -> dict[str, Any]:
    return {"fact_id": str(fact), "page_id": str(page), "quote_verbatim": quote}


async def save_opinion(
    aconn: AsyncConnection, case: uuid.UUID, question: uuid.UUID, body: str, refs: list[dict[str, Any]]
) -> None:
    """Save the opinion on a question as slice E does: sealed text, citations as JSON."""
    sealed = await CRYPTO.encrypt(aconn, case, body.encode(), purpose=opinions.BODY_PURPOSE)
    await aconn.execute(
        "insert into app.opinions (case_id, question_id, body, citation_refs, view) values (%s, %s, %s, %s, 'zz_view')"
        " on conflict (question_id) do update set body = excluded.body, citation_refs = excluded.citation_refs",
        (case, question, sealed, Jsonb(refs)),
    )


def template_headings(data: bytes) -> tuple[list[dict[str, Any]], list[str]]:
    """The heading map slice C reads from a template (paragraph index, text, style), and its style names."""
    document = docx.Document(io.BytesIO(data))
    headings: list[dict[str, Any]] = []
    styles: list[str] = []
    for index, paragraph in enumerate(document.paragraphs):
        style = str(paragraph.style.name) if paragraph.style is not None else ""
        if style and style not in styles:
            styles.append(style)
        if "heading" in style.lower():
            headings.append({"index": index, "text": paragraph.text, "style": style})
    return headings, styles


async def save_template(
    storage: CaseFileStorage,
    admin: psycopg.Connection,
    expert: uuid.UUID,
    mapping: dict[str, int],
    *,
    data: bytes | None = None,
    owner_type: str = "person",
) -> uuid.UUID:
    """Store a template as slice C does (sealed with the cell's box) with the expert's confirmed mapping.

    `expert` is the owner's id: a person, or an organisation when `owner_type` is "organisation"."""
    raw = TEMPLATE.read_bytes()
    headings, styles = template_headings(raw)
    template_id = uuid.uuid4()
    path = f"{'people' if owner_type == 'person' else 'organisations'}/{expert}/template-{template_id}"
    box = SecretBox(CRYPTO.provider)
    stored = data if data is not None else raw
    sealed = box.seal(base64.b64encode(stored).decode(), purpose="report_template", subject=f"{expert}:{template_id}")
    await storage.put(path, sealed.encode())
    admin.execute(
        "insert into app.report_templates (id, owner_type, owner_id, filename, storage_path, sha256, size_bytes,"
        " heading_map, style_map, mapping, version) values (%s, %s, %s, 'template.docx', %s, %s, %s, %s, %s, %s,"
        " (select coalesce(max(version), 0) + 1 from app.report_templates where owner_id = %s))",
        (
            template_id,
            owner_type,
            expert,
            path,
            "a" * 64,
            len(raw),
            Jsonb(headings),
            Jsonb(styles),
            Jsonb([{"index": i, "section": s} for s, i in mapping.items()]),
            expert,
        ),
    )
    return template_id


def template_mapping(combo: Combo) -> tuple[dict[str, int], list[str]]:
    """A confirmed mapping whose template order differs from the pack's section order, and that order."""
    sections = default_library().combine(*combo.codes).report_sections()
    keys = [s["key"] for s in sections]
    indices = [1 + 2 * n for n in range(len(keys))]  # the template's heading paragraphs: 1, 3, 5, ...
    mapping = dict(zip(keys, indices, strict=True))
    first, last = keys[1], keys[-2]
    mapping[first], mapping[last] = mapping[last], mapping[first]
    order = [k for k, _ in sorted(mapping.items(), key=lambda kv: kv[1])]
    assert order != keys, "the template order must differ from the pack order"
    return mapping, order


def app_for(
    api_db: Database, auth_env: dict[str, str], combo: Combo, reads: Reads, storage: CaseFileStorage
) -> FastAPI:
    """The API with storage, case encryption, the real pack library and the stand-in read log."""
    from app.core.auth.supabase import SupabaseAuth
    from app.core.config import auth_settings
    from app.core.email import OutboxMailer
    from app.main import create_app

    return create_app(
        database=api_db,
        region=str(combo.raw("jurisdiction", "pack.yaml")["cell"]),
        auth=SupabaseAuth(auth_settings(auth_env)),
        mailer=OutboxMailer(),
        crypto=CRYPTO,
        storage=storage,
        packs=default_library(),
        page_reads=reads,
    )


def body_paragraphs(data: bytes) -> list[tuple[str, str, str]]:
    """Every paragraph in the document body, inside content controls too: (style name, text, control tag)."""
    document = docx.Document(io.BytesIO(data))
    styles = {s.style_id: s.name for s in document.styles}
    ns = {"w": "http://schemas.openxmlformats.org/wordprocessingml/2006/main"}
    w = "{" + ns["w"] + "}"
    out: list[tuple[str, str, str]] = []
    for p in document.element.body.iter(w + "p"):
        style_el = p.find("w:pPr/w:pStyle", ns)
        style_id = style_el.get(w + "val") if style_el is not None else ""
        name = styles.get(style_id, "Normal") if style_id else "Normal"
        text = "".join(t.text or "" for t in p.iter(w + "t"))
        tag = ""
        parent = p.getparent()
        while parent is not None:
            if parent.tag == w + "sdt":
                tag_el = parent.find("w:sdtPr/w:tag", ns)
                tag = tag_el.get(w + "val") if tag_el is not None else ""
                break
            parent = parent.getparent()
        out.append((str(name), text, tag))
    return out


CLEAN = "The plan was written down [[cite:1]]. The result was low [[cite:2]]."
UNCITED_DATE = "Seen again on 12 March 2024 with no change."


def ready_refs(w: World) -> list[dict[str, Any]]:
    return [cite(w.plan, w.pages[0], "review 3/52"), cite(w.result, w.pages[1], "Ferritin 9")]


async def make_ready(
    seed: Seed,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    w: World,
    reads: Reads,
    *,
    read_all: bool = True,
    decided: bool = True,
    body: str = CLEAN,
) -> uuid.UUID:
    """Everything a report needs before signing, as the other slices leave it: the expert's details,
    conflicts declared, a signature on file, questions locked, every written section saved, the
    checklist ticked and an opinion citing facts on read pages. Returns the case's one finding."""
    from app.features.report.assemble import SECTION_PURPOSE, writable_keys

    pack = default_library().combine(*w.combo.codes)
    if read_all:
        reads.rows.add((w.case, w.pages[2], 3, w.expert))
    admin.execute(
        "insert into app.experts (person_id, registration_body, registration_number) values (%s, 'zz_body', 'ZZ123')"
        " on conflict (person_id) do update set registration_body = 'zz_body', registration_number = 'ZZ123'",
        (w.expert,),
    )
    admin.execute(
        "insert into app.conflict_declarations (case_id, outcome, declared_by) values (%s, 'none', %s)",
        (w.case, w.expert),
    )
    admin.execute(
        "insert into app.signatures (expert_id, method, storage_path, sha256, width, height)"
        " values (%s, 'drawn', %s, %s, 400, 100)",
        (w.expert, f"people/{w.expert}/signature-1", "b" * 64),
    )
    admin.execute(
        "insert into app.instruction_versions (case_id, version_no, kind, received_on)"
        " values (%s, 1, 'letter', '2026-01-05')",
        (w.case,),
    )
    admin.execute("update app.instruction_versions set locked_at = now() where case_id = %s", (w.case,))
    for key in writable_keys(pack):
        sealed = await CRYPTO.encrypt(aconn, w.case, b"Noted as set out.", purpose=SECTION_PURPOSE)
        await aconn.execute(
            "insert into app.report_section_texts (case_id, section, body) values (%s, %s, %s)", (w.case, key, sealed)
        )
    items = [i["key"] for i in pack.compliance_checklist()["items"]]
    admin.execute(
        "insert into app.report_settings (case_id, addressed_contents) values (%s, %s)"
        " on conflict (case_id) do update set addressed_contents = excluded.addressed_contents",
        (w.case, items),
    )
    row = admin.execute(
        "insert into app.findings (case_id, check_key, rule, evidence_fact_ids, evidence_page_ids, pack_version)"
        " values (%s, 'zz_check', 'zz_rule', %s, %s, 'zz@1') returning id",
        (w.case, [w.plan], [w.pages[0]]),
    ).fetchone()
    if decided and row is not None:
        with admin.transaction():
            admin.execute("select set_config('app.person_id', %s, true)", (str(w.expert),))
            admin.execute("update app.findings set decision = 'relevant' where id = %s", (row[0],))
    assert row is not None
    await save_opinion(aconn, w.case, w.question, body, ready_refs(w))
    finding: uuid.UUID = row[0]
    return finding


def signature_png() -> bytes:
    from PIL import Image

    out = io.BytesIO()
    Image.new("RGB", (60, 20), "white").save(out, format="PNG")
    return out.getvalue()


async def store_signature(storage: CaseFileStorage, admin: psycopg.Connection, w: World) -> bytes:
    """Seal the expert's signature image as slice C stores it (cell box, subject person:signature)."""
    import hashlib

    row = admin.execute(
        "select id, storage_path from app.signatures where expert_id = %s order by created_at desc limit 1",
        (w.expert,),
    ).fetchone()
    assert row is not None
    png = signature_png()
    sealed = SecretBox(CRYPTO.provider).seal(
        base64.b64encode(png).decode(), purpose="signature", subject=f"{w.expert}:{row[0]}"
    )
    await storage.put(row[1], sealed.encode())
    admin.execute("update app.signatures set sha256 = %s where id = %s", (hashlib.sha256(png).hexdigest(), row[0]))
    return png
