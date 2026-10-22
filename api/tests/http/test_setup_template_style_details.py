"""Screen 02: each style tile shows what the style looks like, read from the expert's own Word file.

`style_details` lists, for every style the file uses and in the same order as `styles`, its name and the
plain facts the tile shows: bold, italic, point size, numbered, indented. The server reads them; the browser
only words them.
"""

from __future__ import annotations

import io
import uuid

import docx
import pytest
from docx.shared import Cm, Pt

from tests.http.test_setup_template_suggestions import (  # noqa: F401  (pytest fixtures)
    DOCX,
    app,
    expert,
    storage,
)
from tests.support.db import ClientFor

pytestmark = pytest.mark.anyio


def styled_file() -> bytes:
    document = docx.Document()
    heading = document.styles["Heading 1"]
    heading.font.bold = True
    heading.font.size = Pt(14)
    document.add_paragraph("Introduction", style="Heading 1")
    quote = document.styles["Quote"]
    quote.font.italic = True
    quote.paragraph_format.left_indent = Cm(1)
    document.add_paragraph("Said so", style="Quote")
    document.add_paragraph("One", style="List Number")
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


async def test_each_style_carries_its_look(client_for: ClientFor, expert: uuid.UUID) -> None:  # noqa: F811
    client = await client_for(expert)
    up = await client.post(
        "/v1/me/template/file", params={"filename": "t.docx"}, content=styled_file(), headers={"content-type": DOCX}
    )
    assert up.status_code == 201, up.text
    for body in (up.json(), (await client.get("/v1/me/template")).json()):
        details = {d["name"]: d for d in body["style_details"]}
        assert [d["name"] for d in body["style_details"]] == body["styles"]
        assert details["Heading 1"]["bold"] is True and details["Heading 1"]["size_pt"] == 14
        assert details["Quote"]["italic"] is True and details["Quote"]["indented"] is True
        assert details["List Number"]["numbered"] is True
        assert details["Quote"]["bold"] is False and details["Quote"]["size_pt"] is None
