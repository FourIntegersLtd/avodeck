# ruff: noqa: F811  (pytest fixtures imported from the world module)
"""Screen 21b: correct or update a served report (slice G brief, rule 8).

The served version is shown read-only and never changes. A change to one of its paragraphs is a
correction or a change of opinion with a reason; code lists the paragraphs that share its citations, and
T-21e suggests others that may depend on it (the expert reviews each). The supplementary report holds only
the changes, their reasons and the declaration; it links to the served version (`prev_version_id`), needs
step-up to sign, and leaves the served version byte-identical.

Words and paragraphs are short made-up words; nothing here is a real record.
"""

from __future__ import annotations

import hashlib
import io
import uuid
from datetime import timedelta
from typing import Any

import docx
import psycopg
import pytest
from psycopg import AsyncConnection

from app.core.ai.bedrock import ModelRequest
from app.core.db import Database
from app.core.http import etag
from app.core.storage import SupabaseStorage
from tests.fixtures.test_after_serving_world import (  # noqa: F401
    CRYPTO,
    PARAGRAPHS,
    Scripted,
    World,
    aconn,
    app,
    combo,
    mailer,
    model,
    ocr,
    pack_of,
    reads,
    run_job,
    served,
    served_state,
    services,
    storage,
    worker_db,
    world,
)
from tests.support.db import ClientFor

pytestmark = pytest.mark.anyio

NEW_WORDS = "The plan was written down on the first page [1]."
REASON = "I misread the date of the plan."


def change_body(w: World, **over: Any) -> dict[str, Any]:
    body: dict[str, Any] = {
        "paragraph_id": str(w.paragraphs[0]),
        "kind": "correction",
        "new_text": NEW_WORDS,
        "reason": REASON,
        "citation_refs": [{"fact_id": str(w.facts[0]), "page_id": str(w.pages[0]), "quote_verbatim": "review 3/52"}],
    }
    body.update(over)
    return body


def t21e_for(w: World) -> Any:
    def answer(request: ModelRequest) -> dict[str, Any]:
        assert f"[paragraph {w.paragraphs[2]}]" not in request.text, "already found by shared citations"
        assert f"[paragraph {w.paragraphs[1]}]" in request.text and f"[paragraph {w.paragraphs[3]}]" in request.text
        return {
            "items": [
                {"paragraph_id": str(w.paragraphs[3]), "depends": True, "reason": "same visit"},
                {"paragraph_id": str(w.paragraphs[1]), "depends": False, "reason": "other result"},
                {"paragraph_id": str(uuid.uuid4()), "depends": True, "reason": "made up"},
            ]
        }

    return answer


async def run_job_id(worker_db: Database, w: World, job_id: uuid.UUID, kind: str) -> None:
    from app.core.jobs import JobContext, registry

    await registry.kinds[kind].handler(JobContext(worker_db, job_id, w.case, kind, 1))


async def make_change(client_for: ClientFor, worker_db: Database, model: Scripted, w: World) -> dict[str, Any]:
    model.answers["T-21e"] = t21e_for(w)
    resp = await (await client_for(w.expert)).post(w.url("report-changes"), json=change_body(w))
    assert resp.status_code == 201, resp.text
    await run_job_id(worker_db, w, uuid.UUID(resp.json()["job_id"]), "find_dependents")
    got = await (await client_for(w.expert)).get(w.url(f"report-changes/{resp.json()['id']}"))
    assert got.status_code == 200, got.text
    change: dict[str, Any] = got.json()
    return change


async def test_the_served_version_is_shown_read_only(client_for: ClientFor, served: World, world: World) -> None:
    resp = await (await client_for(served.viewer)).get(served.url("report/served"))
    assert resp.status_code == 200, resp.text
    shown = resp.json()
    assert shown["version"]["id"] == str(served.served) and shown["version"]["status"] == "served"
    assert shown["version"]["served_on"] == "2026-03-02" and shown["version"]["sha256"] == "c" * 64
    assert [p["text"] for p in shown["paragraphs"]] == [p[1] for p in PARAGRAPHS]
    assert shown["paragraphs"][2]["fact_ids"] == [str(served.facts[0]), str(served.facts[1])]
    assert (await (await client_for(served.outsider)).get(served.url("report/served"))).status_code == 404
    signed = await (await client_for(served.expert)).post(
        served.url(f"report/versions/{served.served}/sign"), json={"declaration_confirmed": True}
    )
    assert signed.status_code == 409 and signed.json()["code"] == "locked", "the served version never changes"


async def test_a_change_needs_a_reason_and_lists_paragraphs_sharing_its_citations(
    client_for: ClientFor, admin: psycopg.Connection, worker_db: Database, model: Scripted, served: World
) -> None:
    before = served_state(admin, served.served)  # type: ignore[arg-type]
    assert (
        await (await client_for(served.co_expert)).post(served.url("report-changes"), json=change_body(served))
    ).status_code == 403
    assert (
        await (await client_for(served.outsider)).post(served.url("report-changes"), json=change_body(served))
    ).status_code == 404
    expert = await client_for(served.expert)
    for bad in (
        change_body(served, kind="rewrite"),
        change_body(served, reason=""),
        change_body(served, paragraph_id=str(uuid.uuid4())),
        change_body(
            served,
            citation_refs=[
                {"fact_id": str(served.facts[0]), "page_id": str(served.pages[0]), "quote_verbatim": "not there"}
            ],
        ),
    ):
        assert (await expert.post(served.url("report-changes"), json=bad)).status_code == 422
    change = await make_change(client_for, worker_db, model, served)
    assert change["kind"] == "correction" and change["reason"] == REASON
    assert change["old_text"] == PARAGRAPHS[0][1] and change["new_text"] == NEW_WORDS
    assert change["dependent_paragraph_ids"] == [str(served.paragraphs[2])], "shares the cited fact"
    assert change["suggested_dependents"] == [
        {"paragraph_id": str(served.paragraphs[3]), "depends": True, "reason": "same visit"},
        {"paragraph_id": str(served.paragraphs[1]), "depends": False, "reason": "other result"},
    ]
    assert change["supplementary_version_id"] is None
    assert served_state(admin, served.served) == before  # type: ignore[arg-type]


async def test_a_change_is_edited_with_if_match(
    client_for: ClientFor, worker_db: Database, model: Scripted, served: World
) -> None:
    change = await make_change(client_for, worker_db, model, served)
    expert = await client_for(served.expert)
    url = served.url(f"report-changes/{change['id']}")
    assert (await expert.patch(url, json={"reason": "Typed the wrong page."})).status_code == 428
    assert (
        await expert.patch(
            url, json={"reason": "Typed the wrong page."}, headers={"If-Match": etag(change["version"] + 5)}
        )
    ).status_code == 412
    resp = await expert.patch(
        url,
        json={"reason": "Typed the wrong page.", "kind": "change_of_opinion"},
        headers={"If-Match": etag(change["version"])},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["reason"] == "Typed the wrong page." and resp.json()["kind"] == "change_of_opinion"
    listed = (await (await client_for(served.viewer)).get(served.url("report-changes"))).json()
    assert [c["id"] for c in listed["changes"]] == [change["id"]] and listed["supplementary"] is None


async def test_a_supplementary_report_holds_only_the_changes_and_is_signed_with_step_up(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
) -> None:
    before = served_state(admin, served.served)  # type: ignore[arg-type]
    expert = await client_for(served.expert)
    empty = await expert.post(served.url("report-changes/supplementary"))
    assert empty.status_code == 409 and empty.json()["code"] == "no_changes"
    change = await make_change(client_for, worker_db, model, served)
    await run_job(
        worker_db, served, await expert.post(served.url("report-changes/supplementary")), "build_supplementary"
    )
    listed = (await expert.get(served.url("report-changes"))).json()
    draft = listed["supplementary"]
    assert draft["status"] == "draft" and draft["prev_version_id"] == str(served.served)
    assert listed["changes"][0]["supplementary_version_id"] == draft["id"]

    url = served.url(f"report/versions/{draft['id']}/sign")
    assert (await expert.post(url, json={"declaration_confirmed": False})).status_code == 422
    stale = await (await client_for(served.expert, mfa_age=timedelta(minutes=11))).post(
        url, json={"declaration_confirmed": True}
    )
    assert stale.status_code == 403 and stale.json()["code"] == "step_up_required"
    assert (
        await (await client_for(served.co_expert)).post(url, json={"declaration_confirmed": True})
    ).status_code == 403
    signed = await expert.post(url, json={"declaration_confirmed": True})
    assert signed.status_code == 200, signed.text
    version = signed.json()["version"]
    assert version["type"] == "supplementary" and version["status"] == "locked"

    row = admin.execute(
        "select docx_path, prev_version_id from app.report_versions where id = %s", (draft["id"],)
    ).fetchone()
    assert row is not None and row[1] == served.served
    data = await CRYPTO.decrypt(aconn, served.case, await storage.get(row[0]), purpose=row[0])
    assert hashlib.sha256(data).hexdigest() == version["sha256"]
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    assert NEW_WORDS in text and REASON in text and PARAGRAPHS[0][1] in text
    for _, other, _ in PARAGRAPHS[1:]:
        assert other not in text, "only the changes"
    assert pack_of(served.combo).declarations()[0]["text"].split("\n")[0][:40] in text

    assert served_state(admin, served.served) == before, "the served version is byte-identical"  # type: ignore[arg-type]
    late = await expert.patch(
        served.url(f"report-changes/{change['id']}"),
        json={"reason": "Later."},
        headers={"If-Match": etag((await expert.get(served.url(f"report-changes/{change['id']}"))).json()["version"])},
    )
    assert late.status_code == 409, "a change in a signed supplementary report does not change"
    assert (await expert.post(url, json={"declaration_confirmed": True})).status_code == 409
