"""Slice G, screen 20: written questions from the other side.

Acceptance (slice G brief): T-20a on the questions document returns 6 questions as exact substrings; the
planted question asking for a new opinion is suggested as the pack's "goes beyond" classification (the last
one the pack lists) and linked paragraphs are suggested by T-20b; nothing is decided until the expert says so.
"Tell the instructing party" drafts a note with no patient words. Answers cite only `citable_facts()`.
Signing the answers needs step-up and produces a locked version; serving it (slice F's route) serves the
round. Each round of questions has its own deadline and answers document (S11).
"""

# ruff: noqa: F811
from __future__ import annotations

import hashlib
import re
import uuid
from datetime import timedelta
from typing import Any

import psycopg
import pytest
from psycopg import AsyncConnection

from app.core.ai.bedrock import ModelRequest
from app.core.db import Database
from app.core.http import etag
from app.core.storage import SupabaseStorage
from tests.fixtures.test_after_serving_world import (  # noqa: F401
    PAGE_TEXT,
    SERVED_ON,
    Scripted,
    World,
    aconn,
    add_words,
    app,
    classifications,
    combo,
    ids_for,
    mailer,
    model,
    ocr,
    reads,
    run_job,
    served,
    services,
    storage,
    store_document,
    worker_db,
    world,
)
from tests.support.contracts import add_fact
from tests.support.db import ClientFor

pytestmark = pytest.mark.anyio

QUESTIONS = (
    "Please confirm which page records the plan to review in three weeks.",
    "Please state whether the ferritin result was available at the visit.",
    "Please explain what you mean by fell short in paragraph 3.",
    "Please list the pages you relied on for paragraph 2.",
    "Please confirm whether the later visit changes your view in paragraph 4.",
    "Please give a new opinion on whether the outcome would have been different.",
)
DOCUMENT = "Questions to the expert\n" + "\n".join(f"{n}. {q}" for n, q in enumerate(QUESTIONS, start=1))
_LINE = re.compile(r"^\[(question|paragraph) ([0-9a-f-]{36})\] (.*)$", re.M)


def t20a(request: ModelRequest) -> dict[str, Any]:
    return {
        "questions": [
            {"text_verbatim": q, "word_ids": ids_for(request.text, q), "para_ref": None, "reason": "question"}
            for q in QUESTIONS
            if q.split()[2] in request.text
        ]
    }


def t20b_for(kinds: list[str], paragraphs: list[uuid.UUID]) -> Any:
    def answer(request: ModelRequest) -> dict[str, Any]:
        lines = [(m.group(1), m.group(2), m.group(3)) for m in _LINE.finditer(request.text)]
        shown = {uuid.UUID(i) for kind, i, _ in lines if kind == "paragraph"}
        assert shown == set(paragraphs), "every served paragraph is offered for linking"
        items = []
        for kind, qid, text in lines:
            if kind != "question":
                continue
            linked = [str(paragraphs[2])] if "paragraph 3" in text else []
            items.append(
                {
                    "question_id": qid,
                    "kind": kinds[-1] if "new opinion" in text else kinds[0],
                    "paragraph_ids": linked,
                    "reason": "scope",
                }
            )
        return {"items": items}

    return answer


async def open_round(
    client: Any,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    w: World,
    *,
    due: str = "2026-04-20",
    text: str = DOCUMENT,
) -> dict[str, Any]:
    document = await store_document(storage, admin, aconn, w.case, text, "written_questions")
    resp = await client.post(
        w.url("written-question-rounds"),
        json={"received_on": "2026-03-16", "document_id": str(document), "answers_due_on": due},
    )
    assert resp.status_code == 201, resp.text
    made: dict[str, Any] = resp.json()
    return made


async def imported(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    w: World,
    **kw: Any,
) -> tuple[dict[str, Any], list[dict[str, Any]]]:
    model.answers["T-20a"] = t20a
    model.answers["T-20b"] = t20b_for(classifications(w.combo), w.paragraphs)
    client = await client_for(w.expert)
    made = await open_round(client, storage, admin, aconn, w, **kw)
    queued = await client.post(w.url("written-questions/import"), json={"round_id": made["id"]})
    await run_job(worker_db, w, queued, "extract_written_questions")
    listed = await client.get(w.url("written-questions"), params={"round_id": made["id"]})
    assert listed.status_code == 200, listed.text
    return made, listed.json()["questions"]


async def test_a_round_has_its_own_deadline_and_only_the_expert_or_assistant_opens_one(
    client_for: ClientFor, storage: SupabaseStorage, admin: psycopg.Connection, aconn: AsyncConnection, served: World
) -> None:
    document = await store_document(storage, admin, aconn, served.case, DOCUMENT, "written_questions")
    body = {"received_on": "2026-03-16", "document_id": str(document), "answers_due_on": "2026-04-20"}
    for person in (served.co_expert, served.viewer):
        assert (
            await (await client_for(person)).post(served.url("written-question-rounds"), json=body)
        ).status_code == 403
    assert (
        await (await client_for(served.outsider)).post(served.url("written-question-rounds"), json=body)
    ).status_code == 404
    resp = await (await client_for(served.assistant)).post(served.url("written-question-rounds"), json=body)
    assert resp.status_code == 201, resp.text
    made = resp.json()
    assert made["number"] == 1 and made["status"] == "open" and made["answers_version_id"] is None
    deadlines = (await (await client_for(served.expert)).get(served.url("deadlines"))).json()["deadlines"]
    mine = [d for d in deadlines if d["id"] == made["deadline_id"]]
    assert len(mine) == 1 and mine[0]["due_on"] == "2026-04-20" and mine[0]["created_from"] == "round"
    rounds = (await (await client_for(served.viewer)).get(served.url("written-question-rounds"))).json()["rounds"]
    assert [r["id"] for r in rounds] == [made["id"]]


async def test_t20a_finds_the_six_questions_as_exact_words_of_the_document(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
) -> None:
    made, questions = await imported(client_for, storage, admin, aconn, worker_db, model, served)
    assert [q["text_verbatim"] for q in questions] == list(QUESTIONS)
    assert [q["number"] for q in questions] == [1, 2, 3, 4, 5, 6]
    assert [q["para_ref"] for q in questions] == ["1", "2", "3", "4", "5", "6"], "read from the layout by code"
    assert all(q["status"] == "received" and q["round_id"] == made["id"] for q in questions)
    rule = [r for r in served.combo.raw("jurisdiction", "deadlines.yaml")["rules"] if r["key"] == "written_questions"]
    assert all(q["rule_ref"] == rule[0]["rule_ref"] for q in questions)
    for request in model.by_task("T-20a"):
        for words in PAGE_TEXT.values():
            assert words not in request.text, "only the questions document is sent"


async def test_words_the_model_did_not_copy_exactly_are_dropped(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
) -> None:
    def reworded(request: ModelRequest) -> dict[str, Any]:
        good = t20a(request)["questions"]
        bad = dict(good[0], text_verbatim="Please confirm the page with the plan.")
        return {"questions": [bad, *good[1:]]}

    model.answers["T-20a"] = reworded
    model.answers["T-20b"] = t20b_for(classifications(served.combo), served.paragraphs)
    client = await client_for(served.expert)
    made = await open_round(client, storage, admin, aconn, served)
    await run_job(
        worker_db,
        served,
        await client.post(served.url("written-questions/import"), json={"round_id": made["id"]}),
        "extract_written_questions",
    )
    questions = (await client.get(served.url("written-questions"))).json()["questions"]
    assert [q["text_verbatim"] for q in questions] == list(QUESTIONS[1:])


async def test_t20b_suggests_the_beyond_classification_and_links_but_decides_nothing(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
) -> None:
    kinds = classifications(served.combo)
    _, questions = await imported(client_for, storage, admin, aconn, worker_db, model, served)
    by_number = {q["number"]: q for q in questions}
    assert by_number[6]["suggested_classification"] == kinds[-1]
    assert all(by_number[n]["suggested_classification"] == kinds[0] for n in range(1, 6))
    assert by_number[3]["suggested_paragraph_ids"] == [str(served.paragraphs[2])]
    assert all(q["classification"] is None and q["paragraph_ids"] == [] for q in questions), "a suggestion only"
    assert all(q["classification_reason"] == "scope" for q in questions)
    client = await client_for(served.expert)
    q6 = by_number[6]
    resp = await client.patch(
        served.url(f"written-questions/{q6['id']}"),
        json={"classification": kinds[-1]},
        headers={"If-Match": etag(q6["version"])},
    )
    assert resp.status_code == 200, resp.text
    assert resp.json()["classification"] == kinds[-1]
    wrong = await client.patch(
        served.url(f"written-questions/{q6['id']}"),
        json={"classification": "not_a_pack_kind"},
        headers={"If-Match": etag(resp.json()["version"])},
    )
    assert wrong.status_code == 422


async def test_telling_the_instructing_party_drafts_a_note_without_patient_words(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
) -> None:
    _, questions = await imported(client_for, storage, admin, aconn, worker_db, model, served)
    q6 = next(q for q in questions if q["number"] == 6)
    url = served.url(f"written-questions/{q6['id']}/tell-instructing-party")
    assert (await (await client_for(served.co_expert)).post(url)).status_code == 403
    resp = await (await client_for(served.expert)).post(url)
    assert resp.status_code == 200, resp.text
    note = resp.json()
    assert note == {"key": "after_serving.questions.tell_note", "params": {"round": 1, "number": 6, "para_ref": "6"}}
    for words in (*PAGE_TEXT.values(), QUESTIONS[5]):
        for word in words.split():
            if len(word) > 3:
                assert word not in resp.text
    after = (await (await client_for(served.expert)).get(served.url("written-questions"))).json()["questions"]
    assert next(q for q in after if q["id"] == q6["id"])["told_instructing_party_at"] is not None
    audit = admin.execute(
        "select count(*) from app.audit_log where case_id = %s and action = 'written_question.told'", (served.case,)
    ).fetchone()
    assert audit is not None and audit[0] == 1


async def test_answers_cite_only_citable_facts(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
    seed: Any,
) -> None:
    _, questions = await imported(client_for, storage, admin, aconn, worker_db, model, served)
    unread_page = seed.page(served.case, served.records, 40)
    await add_words(admin, aconn, served.case, unread_page, "Iron stopped")
    unread = add_fact(admin, served.case, unread_page, status="verified", value="Iron stopped", word_ids=["w1", "w2"])
    suggested = add_fact(
        admin, served.case, served.pages[1], status="suggested", value="Iron started", word_ids=["w7", "w8"]
    )
    q1 = questions[0]
    client = await client_for(served.expert)
    url = served.url(f"written-questions/{q1['id']}")
    for fact, page in ((unread, unread_page), (suggested, served.pages[1])):
        bad = await client.patch(
            url,
            json={
                "answer_text": "Page 1 [[cite:1]].",
                "citation_refs": [{"fact_id": str(fact), "page_id": str(page), "quote_verbatim": "Iron"}],
            },
            headers={"If-Match": etag(q1["version"])},
        )
        assert bad.status_code == 422, bad.text
    outside = await client.patch(
        url,
        json={
            "answer_text": "Page 1 [[cite:1]].",
            "citation_refs": [
                {"fact_id": str(served.facts[0]), "page_id": str(served.pages[0]), "quote_verbatim": "not in the value"}
            ],
        },
        headers={"If-Match": etag(q1["version"])},
    )
    assert outside.status_code == 422
    assert (
        await (await client_for(served.assistant)).patch(
            url, json={"answer_text": "x"}, headers={"If-Match": etag(q1["version"])}
        )
    ).status_code == 403
    good = await client.patch(
        url,
        json={
            "answer_text": "Page 1 [[cite:1]].",
            "citation_refs": [
                {"fact_id": str(served.facts[0]), "page_id": str(served.pages[0]), "quote_verbatim": "review 3/52"}
            ],
        },
        headers={"If-Match": etag(q1["version"])},
    )
    assert good.status_code == 200, good.text
    assert good.json()["answer_text"] == "Page 1 [[cite:1]]." and good.json()["status"] == "answering"
    assert good.headers["etag"] == etag(good.json()["version"])


async def answer_all(client: Any, w: World, questions: list[dict[str, Any]]) -> None:
    for q in questions:
        resp = await client.patch(
            w.url(f"written-questions/{q['id']}"),
            json={
                "answer_text": f"Answer {q['number']} [[cite:1]].",
                "citation_refs": [
                    {"fact_id": str(w.facts[1]), "page_id": str(w.pages[1]), "quote_verbatim": "Ferritin 9"}
                ],
            },
            headers={"If-Match": etag(q["version"])},
        )
        assert resp.status_code == 200, resp.text


async def test_signing_the_answers_needs_step_up_and_locks_a_new_version(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
) -> None:
    made, questions = await imported(client_for, storage, admin, aconn, worker_db, model, served)
    client = await client_for(served.expert)
    sign = {"round_id": made["id"], "declaration_confirmed": True}
    early = await client.post(served.url("written-questions/answers/sign"), json=sign)
    assert early.status_code == 409 and early.json()["code"] == "answers_missing"
    await answer_all(client, served, questions)
    stale = await (await client_for(served.expert, mfa_age=timedelta(minutes=11))).post(
        served.url("written-questions/answers/sign"), json=sign
    )
    assert stale.status_code == 403 and stale.json()["code"] == "step_up_required"
    assert (
        await (await client_for(served.co_expert)).post(served.url("written-questions/answers/sign"), json=sign)
    ).status_code == 403
    resp = await client.post(served.url("written-questions/answers/sign"), json=sign)
    assert resp.status_code == 200, resp.text
    version = resp.json()["version"]
    assert version["type"] == "answers" and version["status"] == "locked" and version["sha256"]
    row = admin.execute(
        "select status, locked_at, docx_path, prev_version_id from app.report_versions where id = %s", (version["id"],)
    ).fetchone()
    assert row is not None and row[0] == "locked" and row[1] is not None and row[3] == served.served
    snap = admin.execute(
        "select sha256 from app.report_snapshots where report_version_id = %s and kind = 'docx'", (version["id"],)
    ).fetchone()
    assert snap is not None and snap[0] == version["sha256"]
    after = (await client.get(served.url("written-questions"))).json()["questions"]
    assert {q["status"] for q in after} == {"answered"}
    rounds = (await client.get(served.url("written-question-rounds"))).json()["rounds"]
    assert rounds[0]["status"] == "answered" and rounds[0]["answers_version_id"] == version["id"]
    done = [
        d for d in (await client.get(served.url("deadlines"))).json()["deadlines"] if d["id"] == made["deadline_id"]
    ]
    assert done[0]["done_at"] is not None
    changed = await client.patch(
        served.url(f"written-questions/{after[0]['id']}"),
        json={"answer_text": "Changed."},
        headers={"If-Match": etag(after[0]["version"])},
    )
    assert changed.status_code == 409, "a signed answer does not change"
    again = await client.post(served.url("written-questions/answers/sign"), json=sign)
    assert again.status_code == 409


async def test_the_signed_answers_hold_each_question_its_answer_and_the_declaration(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
) -> None:
    import io

    import docx

    from tests.fixtures.test_after_serving_world import CRYPTO, pack_of

    made, questions = await imported(client_for, storage, admin, aconn, worker_db, model, served)
    client = await client_for(served.expert)
    await answer_all(client, served, questions)
    version = (
        await client.post(
            served.url("written-questions/answers/sign"), json={"round_id": made["id"], "declaration_confirmed": True}
        )
    ).json()["version"]
    row = admin.execute("select docx_path from app.report_versions where id = %s", (version["id"],)).fetchone()
    assert row is not None
    data = await CRYPTO.decrypt(aconn, served.case, await storage.get(row[0]), purpose=row[0])
    assert hashlib.sha256(data).hexdigest() == version["sha256"]
    text = "\n".join(p.text for p in docx.Document(io.BytesIO(data)).paragraphs)
    for q in QUESTIONS:
        assert q in text
    assert "Answer 6" in text
    assert pack_of(served.combo).declarations()[0]["text"].split("\n")[0][:40] in text


async def test_serving_the_answers_serves_the_round(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
) -> None:
    made, questions = await imported(client_for, storage, admin, aconn, worker_db, model, served)
    client = await client_for(served.expert)
    await answer_all(client, served, questions)
    version = (
        await client.post(
            served.url("written-questions/answers/sign"), json={"round_id": made["id"], "declaration_confirmed": True}
        )
    ).json()["version"]
    resp = await client.post(
        served.url(f"report/versions/{version['id']}/serve"), json={"served_on": "2026-04-01", "sent_via": "email"}
    )
    assert resp.status_code == 200, resp.text
    assert {q["status"] for q in (await client.get(served.url("written-questions"))).json()["questions"]} == {"served"}
    assert (await client.get(served.url("written-question-rounds"))).json()["rounds"][0]["status"] == "served"
    deadlines = (await client.get(served.url("deadlines"))).json()["deadlines"]
    assert len([d for d in deadlines if d["created_from"] == "pack" and d["kind"] == "written_questions"]) == 1, (
        "serving answers does not count as serving the report again"
    )


async def test_a_second_round_has_its_own_deadline_numbers_and_answers(
    client_for: ClientFor,
    storage: SupabaseStorage,
    admin: psycopg.Connection,
    aconn: AsyncConnection,
    worker_db: Database,
    model: Scripted,
    served: World,
) -> None:
    first, questions = await imported(client_for, storage, admin, aconn, worker_db, model, served)
    second, again = await imported(
        client_for,
        storage,
        admin,
        aconn,
        worker_db,
        model,
        served,
        due="2026-05-11",
        text="Further questions\n1. " + QUESTIONS[1],
    )
    assert second["number"] == 2 and second["deadline_id"] != first["deadline_id"]
    assert [(q["number"], q["text_verbatim"], q["round_id"]) for q in again] == [(1, QUESTIONS[1], second["id"])]
    client = await client_for(served.expert)
    await answer_all(client, served, again)
    signed = await client.post(
        served.url("written-questions/answers/sign"), json={"round_id": second["id"], "declaration_confirmed": True}
    )
    assert signed.status_code == 200, signed.text
    rounds = {r["id"]: r for r in (await client.get(served.url("written-question-rounds"))).json()["rounds"]}
    assert rounds[second["id"]]["status"] == "answered" and rounds[first["id"]]["status"] == "open"
    deadlines = {d["id"]: d for d in (await client.get(served.url("deadlines"))).json()["deadlines"]}
    assert deadlines[second["deadline_id"]]["due_on"] == "2026-05-11"
    assert deadlines[first["deadline_id"]]["done_at"] is None
