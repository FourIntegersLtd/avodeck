"""Signing a draft (screen 18): the signature block added to both files, and what the report is signed with.

`stamp_docx` and `stamp_pdf` add the expert's signature image, name and the date to the end of the draft's
Word file and PDF (the PDF gains one page). `frozen` gathers what the report is signed with (the pack
declarations and the expert's own wording, the signature, the expert's details, the template, the pack
and guideline versions, the instruction versions and the conflict answers) so a later profile change does
not alter what a signed report rests on (architecture 10.6). Snapshot content is sealed under the case key.
"""

from __future__ import annotations

import base64
import hashlib
import io
import json
import uuid
from typing import Any

import docx
import pikepdf
from docx.shared import Inches
from psycopg import AsyncConnection
from reportlab.lib.pagesizes import A4
from reportlab.lib.utils import ImageReader
from reportlab.pdfgen import canvas

from app.core.contracts import profile
from app.core.crypto import KeyProvider, SecretBox
from app.core.packs import MergedPack
from app.features.report.checks import named_expert

SIGNATURE_PURPOSE = "signature"


def snapshot_purpose(kind: str) -> str:
    return f"report_snapshot:{kind}"


def sha256(data: bytes) -> str:
    return hashlib.sha256(data).hexdigest()


def version_hash(docx_sha: str, pdf_sha: str) -> str:
    """The version's fingerprint: SHA-256 over both files' SHA-256 (brief decision)."""
    return sha256(f"docx:{docx_sha}\npdf:{pdf_sha}".encode())


def open_signature(provider: KeyProvider, sealed: bytes, person_id: uuid.UUID, signature_id: uuid.UUID) -> bytes:
    """The signature image as slice C sealed it (cell box, subject person:signature)."""
    opened = SecretBox(provider).open(sealed.decode(), purpose=SIGNATURE_PURPOSE, subject=f"{person_id}:{signature_id}")
    return base64.b64decode(opened)


def stamp_docx(data: bytes, image: bytes, lines: list[str]) -> bytes:
    document = docx.Document(io.BytesIO(data))
    document.add_paragraph().add_run().add_picture(io.BytesIO(image), width=Inches(2.5))
    for line in lines:
        document.add_paragraph(line)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def stamp_pdf(data: bytes, image: bytes, lines: list[str]) -> bytes:
    page = io.BytesIO()
    sheet = canvas.Canvas(page, pagesize=A4)
    width, height = A4
    reader = ImageReader(io.BytesIO(image))
    w, h = reader.getSize()
    shown = min(180.0, float(w))
    sheet.drawImage(reader, 72, height - 72 - shown * h / w, width=shown, height=shown * h / w)
    y = height - 100 - shown * h / w
    for line in lines:
        sheet.drawString(72, y, line)
        y -= 18
    sheet.showPage()
    sheet.save()
    with pikepdf.open(io.BytesIO(data)) as report, pikepdf.open(io.BytesIO(page.getvalue())) as signed:
        report.pages.extend(signed.pages)
        out = io.BytesIO()
        report.save(out)
    return out.getvalue()


def canonical(content: Any) -> bytes:
    return json.dumps(content, sort_keys=True, separators=(",", ":"), default=str).encode()


async def frozen(
    conn: AsyncConnection, case_id: uuid.UUID, pack: MergedPack, signature: profile.Signature
) -> dict[str, Any]:
    """What the report is signed with, by snapshot kind."""
    expert = await named_expert(conn, case_id) or {}
    template = await profile.report_template(conn, case_id)
    conflicts = await profile.conflict_declaration(conn, case_id)
    cur = await conn.execute(
        "select version_no, kind, received_on, locked_at from app.instruction_versions"
        " where case_id = %s and locked_at is not null order by version_no",
        (case_id,),
    )
    instructions = [
        {"version_no": v, "kind": k, "received_on": r, "locked_at": at} for v, k, r, at in await cur.fetchall()
    ]
    return {
        "declarations": {
            "pack": [
                {"key": d["key"], "rule_ref": d.get("rule_ref"), "version": d.get("version"), "text": d["text"]}
                for d in pack.declarations()
            ],
            "added": [
                {"kind": d.kind, "text": d.text, "locked_at": d.locked_at}
                for d in await profile.declarations(conn, case_id)
            ],
        },
        "signature": {"id": signature.id, "sha256": signature.sha256, "method": signature.method},
        "expert": {k: expert.get(k) for k in ("person_id", "full_name", "registration_body", "registration_number")},
        "template": (
            {"id": template.id, "version": template.version, "sha256": template.sha256}
            if template is not None
            else None
        ),
        "packs": pack.versions,
        "guidelines": pack.guidelines(),
        "instructions": instructions,
        "conflicts": (
            {
                "outcome": conflicts.outcome,
                "declared_at": conflicts.declared_at,
                "answers": [
                    {"question_key": a.question_key, "answer": a.answer, "has_detail": a.detail is not None}
                    for a in conflicts.answers
                ],
            }
            if conflicts is not None
            else None
        ),
    }
