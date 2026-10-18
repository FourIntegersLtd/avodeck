"""Lay the assembled report out in the expert's own Word template (python-docx).

Paragraphs before the first mapped heading (letterhead, title) are kept. Each section follows in the
order `assemble` decided: a mapped section keeps the template's own heading paragraph (its words and
style, unchanged); any other section gets a copy of a template heading with the section's words. The
section's paragraphs sit in one block-level content control tagged `section:<key>`, in the template's
body style (the style of the paragraph after that heading). The declarations control is locked.

F16: if the template cannot be opened or its mapping does not fit it, the report is laid out in the
standard layout instead and `fallback` says so.
"""

from __future__ import annotations

import copy
import io
from collections.abc import Callable
from typing import Any

import docx
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

from app.core.contracts.profile import ReportTemplate
from app.features.report.assemble import Assembly, Block, template_mapping
from app.features.report.words import section_label

Element = Any


def _paragraph(text: str, style_id: str | None) -> Element:
    p = OxmlElement("w:p")
    if style_id:
        ppr = OxmlElement("w:pPr")
        ps = OxmlElement("w:pStyle")
        ps.set(qn("w:val"), style_id)
        ppr.append(ps)
        p.append(ppr)
    if text:
        r = OxmlElement("w:r")
        t = OxmlElement("w:t")
        t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
        t.text = text
        r.append(t)
        p.append(r)
    return p


def _style_id(p: Element) -> str | None:
    ps = p.find(qn("w:pPr") + "/" + qn("w:pStyle"))
    return None if ps is None else str(ps.get(qn("w:val")))


def _relabel(heading: Element, text: str) -> Element:
    out = copy.deepcopy(heading)
    runs = out.findall(qn("w:r"))
    rpr = runs[0].find(qn("w:rPr")) if runs else None
    for child in list(out):
        if child.tag != qn("w:pPr"):
            out.remove(child)
    r = OxmlElement("w:r")
    if rpr is not None:
        r.append(copy.deepcopy(rpr))
    t = OxmlElement("w:t")
    t.set("{http://www.w3.org/XML/1998/namespace}space", "preserve")
    t.text = text
    r.append(t)
    out.append(r)
    return out


def _control(block: Block, style_id: str | None, control_id: int) -> Element:
    sdt = OxmlElement("w:sdt")
    pr = OxmlElement("w:sdtPr")
    for name, value in (("w:alias", block.key), ("w:tag", f"section:{block.key}"), ("w:id", str(control_id))):
        el = OxmlElement(name)
        el.set(qn("w:val"), value)
        pr.append(el)
    if block.locked:
        lock = OxmlElement("w:lock")
        lock.set(qn("w:val"), "sdtContentLocked")
        pr.append(lock)
    sdt.append(pr)
    content = OxmlElement("w:sdtContent")
    for para in block.paragraphs or []:
        content.append(_paragraph(para.text, style_id))
    if not block.paragraphs:
        content.append(_paragraph("", style_id))
    sdt.append(content)
    return sdt


def _write(
    document: Any,
    assembly: Assembly,
    heading_for: Callable[[str], Element],
    body_style_for: Callable[[str], str | None],
) -> bytes:
    body = document.element.body
    sect = body.find(qn("w:sectPr"))
    for n, block in enumerate(assembly.blocks, start=1):
        for el in (heading_for(block.key), _control(block, body_style_for(block.key), 1000 + n)):
            if sect is not None:
                sect.addprevious(el)
            else:
                body.append(el)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def _in_template(assembly: Assembly, data: bytes, template: ReportTemplate) -> bytes:
    document = docx.Document(io.BytesIO(data))
    paragraphs = document.paragraphs
    keys = [b.key for b in assembly.blocks]
    mapping = template_mapping(keys, template)
    if not mapping or any(i < 0 or i >= len(paragraphs) for i in mapping.values()):
        raise ValueError("mapping does not fit the template")
    headings = {int(h["index"]) for h in template.heading_map if isinstance(h, dict) and "index" in h}
    headings |= set(mapping.values())
    elements = [p._p for p in paragraphs]
    plain = [_style_id(el) for i, el in enumerate(elements) if i not in headings and _style_id(el)]
    default_body = plain[0] if plain else None

    def body_after(index: int) -> str | None:
        for j in range(index + 1, len(elements)):
            if j in headings:
                break
            if _style_id(elements[j]):
                return _style_id(elements[j])
        return default_body

    kept = {k: copy.deepcopy(elements[i]) for k, i in mapping.items()}
    styles = {k: body_after(i) for k, i in mapping.items()}
    sample = kept[min(mapping, key=lambda k: mapping[k])]
    body = document.element.body
    first = elements[min(mapping.values())]
    children = list(body)
    for child in children[children.index(first) :]:
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    return _write(
        document,
        assembly,
        lambda key: kept[key] if key in kept else _relabel(sample, section_label(key)),
        lambda key: styles.get(key, default_body),
    )


def _standard(assembly: Assembly) -> bytes:
    document = docx.Document()
    body = document.element.body
    for child in list(body):
        if child.tag != qn("w:sectPr"):
            body.remove(child)
    heading = document.styles["Heading 1"].style_id
    normal = document.styles["Normal"].style_id
    return _write(document, assembly, lambda key: _paragraph(section_label(key), heading), lambda key: normal)


def render_docx(assembly: Assembly, data: bytes | None, template: ReportTemplate | None) -> tuple[bytes, bool]:
    """The Word file, and whether it fell back to the standard layout (F16)."""
    if data is not None and template is not None:
        try:
            return _in_template(assembly, data, template), False
        except Exception:  # noqa: BLE001 - any template problem falls back (F16); no content is logged
            pass
    return _standard(assembly), True
