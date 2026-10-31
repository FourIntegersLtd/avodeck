"""Slice B job: `export_transcript` builds the Word copy of the records transcript (ADR 0028).

The file holds the entries as the screen does, in the same order, with names and dates of birth
stripped (rule 7). It is stored sealed under the case key and handed over once, to the person who asked
(the route in `transcript_routes`); slice H's daily run removes a copy nobody took (`EXPORT_KEEP_DAYS`).
Logs carry exception types only and never a word of the record (rule 6).
"""

from __future__ import annotations

import io
import uuid

import docx
from psycopg import AsyncConnection

from app.core import audit
from app.core.crypto import CaseCrypto, key_provider_from_env
from app.core.errors import ApiError
from app.core.jobs import JobContext, JobRegistry, registry
from app.core.packs import PackLibrary, default_library
from app.core.storage import StorageError, SupabaseStorage, case_prefix, storage_settings
from app.features.reader import repo, transcript
from app.features.reader.transcript_words import label_text, show_date, word

KIND = "export_transcript"
PAGE = 200
DOWNLOADED = "case.transcript_downloaded"


class _Services:
    """Storage, the case keys and the packs for the job, made from the cell config on first use."""

    def __init__(self) -> None:
        self.storage: SupabaseStorage | None = None
        self.crypto: CaseCrypto | None = None
        self.packs: PackLibrary | None = None

    def get_storage(self) -> SupabaseStorage:
        if self.storage is None:
            self.storage = SupabaseStorage(storage_settings())
        return self.storage

    def get_crypto(self) -> CaseCrypto:
        if self.crypto is None:
            self.crypto = CaseCrypto(key_provider_from_env())
        return self.crypto

    def get_packs(self) -> PackLibrary:
        if self.packs is None:
            self.packs = default_library()
        return self.packs


services = _Services()


def transcript_path(case_id: uuid.UUID, job_id: uuid.UUID) -> str:
    return f"{case_prefix(case_id)}/exports/{job_id}/transcript.docx"


def build(
    entries: list[transcript.TranscriptEntry], *, date_order: str, language: str, unread_pages: list[int] | None = None
) -> bytes:
    document = docx.Document()
    document.add_heading(word(language, "title"), level=0)
    for e in entries:
        parts = [word(language, "page", number=e.page_no), show_date(e.event_date, date_order, language)]
        parts.append(label_text(e.kind_label))
        if e.author_label:
            parts.append(label_text(e.author_label))
        document.add_heading(", ".join(parts), level=2)
        for line in e.lines:
            document.add_paragraph(line)
    for number in unread_pages or []:
        document.add_heading(word(language, "page", number=number), level=2)
        document.add_paragraph(word(language, "unread"))
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


async def export_transcript(ctx: JobContext) -> None:
    """Builds the requester's copy. A job nobody asked for (`requested_by` empty) builds nothing."""
    storage, crypto, library = services.get_storage(), services.get_crypto(), services.get_packs()
    async with ctx.transaction() as conn:
        cur = await conn.execute(
            "select 1 from app.jobs where id = %s and requested_by is not null and exists"
            " (select 1 from app.cases where id = %s and status <> 'deleted')",
            (ctx.job_id, ctx.case_id),
        )
        if await cur.fetchone() is None:
            return
        cur = await conn.execute("select named_expert_id from app.cases where id = %s", (ctx.case_id,))
        named = (await cur.fetchone() or (None,))[0]
        codes = await repo.case_codes(conn, ctx.case_id)
        if named is None or codes is None:
            return
        pack = library.for_case(codes)
        everything: list[transcript.TranscriptEntry] = []
        unread: list[int] = []
        while True:
            got = await transcript.entries(
                conn, crypto, library, ctx.case_id, offset=len(everything), limit=PAGE, only_read_by=named
            )
            everything += got.entries
            unread = got.unread_pages
            if not got.entries or len(everything) >= got.total:
                break
        path = transcript_path(ctx.case_id, ctx.job_id)
        document = build(everything, date_order=pack.date_order, language=pack.language, unread_pages=unread)
        await storage.put(path, await crypto.encrypt(conn, ctx.case_id, document, purpose=path))


def _missing(error: StorageError) -> bool:
    """True when storage said the object is not there (not any other fault)."""
    return str(error).rsplit(" ", 1)[-1] in ("400", "404")


async def take(
    conn: AsyncConnection, storage: SupabaseStorage, crypto: CaseCrypto, case_id: uuid.UUID, job_id: uuid.UUID
) -> bytes:
    """The Word file, once. That it was taken is the audit row, so a copy that could not be removed is never
    handed over twice; the caller removes the stored copy after the transaction commits (`remove`)."""
    path = transcript_path(case_id, job_id)
    await conn.execute("select pg_advisory_xact_lock(hashtextextended(%s, 0))", (path,))
    cur = await conn.execute(
        "select 1 from app.audit_log where case_id = %s and action = %s and target = %s",
        (case_id, DOWNLOADED, f"export:{job_id}"),
    )
    if await cur.fetchone() is not None:
        raise ApiError("not_found", "This file has already been downloaded.", 404)
    try:
        blob = await storage.get(path)
    except StorageError as error:
        if _missing(error):
            raise ApiError("not_found", "This file has already been downloaded.", 404) from None
        raise ApiError("unavailable", "File storage is not ready.", 503) from None
    data = await crypto.decrypt(conn, case_id, blob, purpose=path)
    await audit.record(conn, action=DOWNLOADED, target=f"export:{job_id}", case_id=case_id)
    return data


async def remove(storage: SupabaseStorage, case_id: uuid.UUID, job_id: uuid.UUID) -> None:
    """Removes the stored copy once the download is recorded. A failure leaves it for the daily sweep."""
    try:
        await storage.delete([transcript_path(case_id, job_id)])
    except StorageError:
        pass


def register(reg: JobRegistry) -> None:
    reg.kind(KIND, changes_case=False, runs_on_hold=True)(export_transcript)


register(registry)
