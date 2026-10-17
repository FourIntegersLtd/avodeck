"""Files that are not PDFs (F36, architecture 10.3 "Inputs other than PDF").

`unpack` sniffs one file by its bytes (the filename only tells an email apart) and returns its own
pages as one PDF, the files inside it (each becomes a document of its own) and whether it is an
exhibit. Images become image pages, read by OCR as usual. Word files and emails are typed onto
pages whose words come from the file itself. Audio, video, medical imaging, anything unknown and
archives too large to unpack are exhibits: kept, listed, never paged. Nothing here is logged.
"""

from __future__ import annotations

import io
import re
import textwrap
import zipfile
from dataclasses import dataclass, field
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from html.parser import HTMLParser

import pikepdf
from PIL import Image, ImageSequence

PDF = "application/pdf"
ZIP = "application/zip"
EMAIL = "message/rfc822"
DOCX = "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
UNKNOWN = "application/octet-stream"

PAGED = {PDF, EMAIL, DOCX, "image/jpeg", "image/png", "image/tiff"}

A4 = (595, 842)
MARGIN = 72
FONT_SIZE = 11
LEADING = 14
LINE_CHARS = 84
IMAGE_DPI = 200

_HEADER = re.compile(rb"^(From|To|Subject|Date|MIME-Version|Received|Return-Path|Message-ID):", re.I | re.M)


@dataclass(frozen=True)
class Limits:
    max_bytes: int = 500 * 1024 * 1024  # what a zip may unpack to
    max_members: int = 2000
    max_depth: int = 3  # zips and emails inside each other


@dataclass(frozen=True)
class Member:
    filename: str
    data: bytes = field(repr=False)


@dataclass(frozen=True)
class Unpacked:
    media_type: str
    pdf: bytes | None = field(default=None, repr=False)
    members: list[Member] = field(default_factory=list)
    exhibit: bool = False


def sniff(data: bytes, filename: str = "") -> str:
    head = data[:64]
    if b"%PDF-" in data[:1024]:
        return PDF
    if head.startswith(b"\xff\xd8\xff"):
        return "image/jpeg"
    if head.startswith(b"\x89PNG\r\n\x1a\n"):
        return "image/png"
    if head[:4] in (b"II*\x00", b"MM\x00*"):
        return "image/tiff"
    if head.startswith(b"PK\x03\x04"):
        try:
            with zipfile.ZipFile(io.BytesIO(data)) as z:
                return DOCX if "word/document.xml" in z.namelist() else ZIP
        except zipfile.BadZipFile:
            return UNKNOWN
    if head.startswith(b"RIFF") and head[8:12] == b"WAVE":
        return "audio/wav"
    if head.startswith(b"RIFF") and head[8:12] == b"AVI ":
        return "video/x-msvideo"
    if head.startswith((b"ID3", b"\xff\xfb", b"\xff\xf3")):
        return "audio/mpeg"
    if head.startswith(b"OggS"):
        return "audio/ogg"
    if head.startswith(b"fLaC"):
        return "audio/flac"
    if head.startswith(b"\x1aE\xdf\xa3"):
        return "video/x-matroska"
    if head[4:8] == b"ftyp":
        brand = head[8:12]
        if brand in (b"heic", b"heix", b"mif1", b"heif"):
            return "image/heic"
        if brand.startswith(b"M4A"):
            return "audio/mp4"
        if brand == b"qt  ":
            return "video/quicktime"
        return "video/mp4"
    if data[128:132] == b"DICM":
        return "application/dicom"
    if filename.lower().endswith(".eml") or _HEADER.match(data[:2048]):
        return EMAIL
    return UNKNOWN


def unpack(data: bytes, filename: str, limits: Limits | None = None) -> Unpacked:
    limits = limits or Limits()
    media = sniff(data, filename)
    if media == PDF:
        return Unpacked(media, pdf=data)
    if media == ZIP:
        members = _zip_members(data, limits)
        return Unpacked(media, members=members) if members is not None else Unpacked(media, exhibit=True)
    if media not in PAGED:
        return Unpacked(media, exhibit=True)
    try:
        if media == DOCX:
            return Unpacked(media, pdf=_typed(_docx_lines(data)))
        if media == EMAIL:
            lines, members = _email(data)
            return Unpacked(media, pdf=_typed(lines), members=members)
        return Unpacked(media, pdf=_image_pages(data))
    except Exception:  # noqa: BLE001 - a file we cannot turn into pages is kept as an exhibit
        return Unpacked(media, exhibit=True)


# --- archives ------------------------------------------------------------------------------------


def _zip_members(data: bytes, limits: Limits) -> list[Member] | None:
    try:
        with zipfile.ZipFile(io.BytesIO(data)) as z:
            infos = [i for i in z.infolist() if not i.is_dir() and not i.filename.startswith("__MACOSX/")]
            if len(infos) > limits.max_members or sum(i.file_size for i in infos) > limits.max_bytes:
                return None
            if any(i.flag_bits & 0x1 for i in infos):
                return None  # locked with a password
            return sorted((Member(i.filename, z.read(i)) for i in infos), key=lambda m: m.filename)
    except (zipfile.BadZipFile, zipfile.LargeZipFile, RuntimeError, ValueError):
        return None


# --- emails --------------------------------------------------------------------------------------


class _Text(HTMLParser):
    def __init__(self) -> None:
        super().__init__()
        self.parts: list[str] = []

    def handle_data(self, data: str) -> None:
        self.parts.append(data)

    def handle_starttag(self, tag: str, attrs: object) -> None:
        if tag in ("br", "p", "div", "tr", "li"):
            self.parts.append("\n")


def _email(data: bytes) -> tuple[list[str], list[Member]]:
    message = BytesParser(policy=policy.default).parsebytes(data)
    assert isinstance(message, EmailMessage)
    lines = [f"{h}: {message[h]}" for h in ("From", "To", "Cc", "Date", "Subject") if message[h]]
    lines.append("")
    body = message.get_body(preferencelist=("plain", "html"))
    if body is not None:
        text = str(body.get_content())
        if body.get_content_type() == "text/html":
            reader = _Text()
            reader.feed(text)
            text = "".join(reader.parts)
        lines.extend(text.splitlines())
    members: list[Member] = []
    for n, part in enumerate(message.iter_attachments(), start=1):
        name = part.get_filename() or f"attachment-{n}"
        if part.get_content_type() == EMAIL:
            inner = part.get_content()
            payload = inner.as_bytes() if isinstance(inner, EmailMessage) else bytes(inner)
        else:
            raw = part.get_payload(decode=True)
            payload = raw if isinstance(raw, bytes) else b""
        members.append(Member(name, bytes(payload)))
    return lines, members


# --- Word files ----------------------------------------------------------------------------------


def _docx_lines(data: bytes) -> list[str]:
    from docx import Document

    document = Document(io.BytesIO(data))
    lines = [p.text for p in document.paragraphs]
    for table in document.tables:
        for row in table.rows:
            lines.append("  ".join(cell.text for cell in row.cells))
    return lines


# --- writing pages -------------------------------------------------------------------------------


def _save(pdf: pikepdf.Pdf) -> bytes:
    out = io.BytesIO()
    pdf.save(out, deterministic_id=True)
    return out.getvalue()


def _escape(line: str) -> bytes:
    raw = line.encode("cp1252", errors="replace")
    return raw.replace(b"\\", b"\\\\").replace(b"(", b"\\(").replace(b")", b"\\)")


def _typed(lines: list[str]) -> bytes:
    """Text typed in Helvetica onto A4 pages, wrapped, so its words come from the text layer."""
    wrapped: list[str] = []
    for line in lines:
        clean = line.replace("\t", "    ").rstrip()
        wrapped.extend(textwrap.wrap(clean, LINE_CHARS) or [""])
    per_page = (A4[1] - 2 * MARGIN) // LEADING
    chunks = [wrapped[i : i + per_page] for i in range(0, len(wrapped), per_page)] or [[]]
    pdf = pikepdf.new()
    font = pdf.make_indirect(
        pikepdf.Dictionary(
            Type=pikepdf.Name.Font,
            Subtype=pikepdf.Name.Type1,
            BaseFont=pikepdf.Name.Helvetica,
            Encoding=pikepdf.Name.WinAnsiEncoding,
        )
    )
    for chunk in chunks:
        ops = b"".join(
            b"BT /F1 %d Tf %d %d Td (%s) Tj ET\n" % (FONT_SIZE, MARGIN, A4[1] - MARGIN - LEADING * i, _escape(t))
            for i, t in enumerate(chunk)
            if t.strip()
        )
        pdf.pages.append(
            pikepdf.Page(
                pikepdf.Dictionary(
                    Type=pikepdf.Name.Page,
                    MediaBox=[0, 0, *A4],
                    Resources=pikepdf.Dictionary(Font=pikepdf.Dictionary(F1=font)),
                    Contents=pdf.make_stream(ops),
                )
            )
        )
    return _save(pdf)


def _image_pages(data: bytes) -> bytes:
    """One image-only page per image frame, for OCR to read."""
    with Image.open(io.BytesIO(data)) as image:
        frames = [frame.convert("RGB") for frame in ImageSequence.Iterator(image)]
    out = io.BytesIO()
    first, *rest = frames
    first.save(out, "PDF", save_all=True, append_images=rest, resolution=IMAGE_DPI)
    with pikepdf.open(io.BytesIO(out.getvalue())) as pdf:
        pdf.docinfo = pdf.make_indirect(pikepdf.Dictionary())
        return _save(pdf)
