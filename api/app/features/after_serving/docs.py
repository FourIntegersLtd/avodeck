"""Slice G's Word documents: the answers to written questions and the supplementary report.

A document is a list of (style, text) blocks; signing adds the dated line to the same file. Each block is
one paragraph, so the declaration's lines stay as the pack writes them.
"""

from __future__ import annotations

import io
from collections.abc import Sequence

import docx

Block = tuple[str, str]  # (style: "title" | "heading" | "body", text)


def build(blocks: Sequence[Block]) -> bytes:
    document = docx.Document()
    for style, text in blocks:
        if style == "title":
            document.add_heading(text, level=0)
        elif style == "heading":
            document.add_heading(text, level=1)
        else:
            document.add_paragraph(text)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def append(data: bytes, lines: Sequence[str]) -> bytes:
    document = docx.Document(io.BytesIO(data))
    for line in lines:
        document.add_paragraph(line)
    out = io.BytesIO()
    document.save(out)
    return out.getvalue()


def declaration_blocks(declarations: Sequence[dict[str, str]], heading: str) -> list[Block]:
    blocks: list[Block] = [("heading", heading)]
    for declaration in declarations:
        blocks.extend(("body", line) for line in str(declaration.get("text", "")).split("\n") if line.strip())
    return blocks
