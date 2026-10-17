"""The other synthetic documents for Mrs A, built after records.pdf.

Usage: python evals/generator/extras.py OUT_DIR   (OUT_DIR as used by build.py)
"""
from __future__ import annotations

import io
import json
import os
import random
import sys
import zipfile
from datetime import date, datetime

from pypdf import PdfReader, PdfWriter
from reportlab.lib.units import mm
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import A4

sys.path.insert(0, os.path.dirname(__file__))

from common import Doc, W, H, uk_date, long_date, fmt_nhs  # noqa: E402
from story import PATIENT  # noqa: E402
from scan import pdf_pages_to_images, degrade, images_to_pdf  # noqa: E402

P = PATIENT
FIRM = ("Harland & Price LLP", ["Solicitors", "9 Cathedral Walk", "Northam NM1 1RP", "Ref: HP/CN/2291/RT"])


def letterhead(d: Doc, firm=FIRM):
    c = d.c
    c.setFont("Serif-B", 15)
    c.drawString(d.ml, H - 22 * mm, firm[0])
    c.setFont("Serif", 8.5)
    for i, ln in enumerate(firm[1]):
        c.drawRightString(W - d.mr, H - 18 * mm - i * 4 * mm, ln)
    d.y = H - 42 * mm


QUESTIONS = [
    "Please set out your view of the standard of care provided by the general practitioners between February and August 2024.",
    "In particular, did the standard of care provided at the consultation on 15 May 2024 fall below that of a reasonably competent general practitioner, and if so in what respects?",
    "Would a reasonably competent general practitioner have referred Mrs Ashworth on the urgent suspected cancer pathway, and if so by what date?",
    "If Mrs Ashworth had been referred earlier, on the balance of probabilities would her cancer have been diagnosed at an earlier stage?",
]
ASSUMPTION = "You should assume that Mrs Ashworth's account of her symptoms in her witness statement is accurate unless the records clearly show otherwise."


def letter_of_instruction(out, exp=None):
    d = Doc("letter-of-instruction", "loi", margin_l=24 * mm, margin_r=24 * mm)
    letterhead(d)
    for ln in ["Dr A Expert", "Expert Witness (General Practice)", "By secure email"]:
        d.text(d.ml, ln, "Serif", 10.5)
    d.y -= 8
    d.text(d.ml, long_date(date(2025, 10, 9)), "Serif", 10.5, lead=18)
    d.text(d.ml, "Dear Dr Expert", "Serif", 10.5, lead=16)
    d.text(d.ml, "Margaret Ashworth v Brookvale Surgery (Dr R Patel and Dr S Okafor)", "Serif-B", 10.5)
    d.text(d.ml, "Letter of instruction - breach of duty, general practice", "Serif-B", 10.5, lead=18)
    paras = [
        "1. We act for the Claimant, Mrs Margaret Ashworth (date of birth 14/03/1969), in a claim against the general "
        "practitioners at Brookvale Surgery. Mrs Ashworth was diagnosed with adenocarcinoma of the sigmoid colon in "
        "September 2024. She alleges that the diagnosis was delayed.",
        "2. We enclose the paginated medical records (450 pages), the Claimant's witness statement and a copy of Part 35 "
        "and Practice Direction 35. Please confirm you have no conflict of interest before reading the records.",
        "3. Your report should comply with Part 35 and should be addressed to the court. We should be grateful for your report "
        "within 8 weeks. Please tell us immediately if you need further records.",
        "4. Questions. Please answer the following questions:",
    ]
    for p in paras:
        d.para(d.ml, p, W - d.ml - d.mr, "Serif", 10.5, lead=14)
        d.y -= 6
    for i, q in enumerate(QUESTIONS, 1):
        d.para(d.ml + 8 * mm, f"({chr(96 + i)}) {q}", W - d.ml - d.mr - 8 * mm, "Serif", 10.5, lead=14)
        d.y -= 4
    d.y -= 4
    d.para(d.ml, f"5. Assumptions. {ASSUMPTION}", W - d.ml - d.mr, "Serif", 10.5, lead=14)
    d.y -= 6
    d.para(d.ml, "6. We understand question (d) concerns causation and may fall outside your expertise. If so, please say so "
                 "and we will instruct a colorectal surgeon.", W - d.ml - d.mr, "Serif", 10.5, lead=14)
    d.y -= 14
    d.text(d.ml, "Yours sincerely", "Serif", 10.5, lead=26)
    d.text(d.ml, "R. Tennant, Partner, Harland & Price LLP", "Serif", 10.5)
    p = d.save(out)
    json.dump({"questions": [{"number": i, "text_verbatim": q, "para_ref": f"4({chr(96 + i)})",
                              "kind": "causation" if i == 4 else "breach", "outside_expertise": i == 4}
                             for i, q in enumerate(QUESTIONS, 1)],
               "assumptions": [{"text_verbatim": ASSUMPTION, "para_ref": "5"}]},
              open(os.path.join(exp or EXP, "questions.json"), "w"), indent=1)
    return p


def witness_statement(out):
    d = Doc("witness-statement", "witness", margin_l=26 * mm, margin_r=26 * mm)
    c = d.c
    c.setFont("Serif", 9)
    c.drawRightString(W - d.mr, H - 18 * mm, "Claimant / M Ashworth / 1st / Exhibits: none / 02.10.2025")
    d.y = H - 34 * mm
    d.text(d.ml, "IN THE COUNTY COURT AT NORTHAM", "Serif-B", 10.5)
    d.text(d.ml, "BETWEEN: MARGARET ASHWORTH (Claimant) and DR RAKESH PATEL and DR SOLA OKAFOR (Defendants)", "Serif", 9.5, lead=18)
    d.text(d.ml, "WITNESS STATEMENT OF MARGARET ASHWORTH", "Serif-B", 11, lead=20)
    paras = [
        "I, Margaret Ashworth, of 27 Larch Close, Northam, will say as follows:",
        "1. In early 2024 I felt tired all the time. I saw Dr Patel in February and he arranged blood tests. I was not told "
        "that the results were abnormal, although later a doctor mentioned I was a bit anaemic.",
        "2. From about February my bowels changed. I had loose stools most days and cramping pains. I told Dr Patel about "
        "this on the telephone in March.",
        "3. I saw Dr Okafor on 15 May 2024. I told her I was still exhausted, that my bowels were still loose and that I had "
        "lost weight because my clothes were loose. She started me on iron tablets. I was never told to come back. I "
        "thought the iron tablets would sort it out.",
        "4. When I rang the surgery on 20 May about constipation from the iron, the nurse told me to take them with food. "
        "I do not remember being told to call back about anything else.",
        "5. In June I saw Dr Patel again and he said he would refer me to the hospital. I heard nothing from the hospital "
        "for weeks. In August I saw blood in my stools and Dr Okafor referred me urgently.",
        "6. I was diagnosed with bowel cancer in September 2024 and had an operation in October followed by chemotherapy.",
        "I believe that the facts stated in this witness statement are true. I understand that proceedings for contempt of "
        "court may be brought against anyone who makes, or causes to be made, a false statement in a document verified by a "
        "statement of truth without an honest belief in its truth.",
    ]
    for p in paras:
        d.para(d.ml, p, W - d.ml - d.mr, "Serif", 10.5, lead=15)
        d.y -= 7
    d.text(d.ml, "Signed: M Ashworth          Dated: 2 October 2025", "Serif", 10.5)
    return d.save(out)


def written_questions(out):
    qs = [
        "Please confirm the page reference for your statement that the ferritin result of February 2024 was not acted upon.",
        "At paragraph 6.3 you refer to the plan of 15 May 2024. Please confirm the exact wording recorded in the notes.",
        "Do you accept that iron replacement was a reasonable initial step on 15 May 2024?",
        "Please clarify whether your opinion on the date by which referral should have been made relies on the haematology letter.",
        "Please explain what you mean by 'safety-netting was inadequate' at paragraph 7.1.",
        "Please now give your opinion on whether, with earlier referral, the tumour would have been at an earlier stage and the "
        "Claimant's prognosis better.",
    ]
    d = Doc("written-questions", "written_questions", margin_l=24 * mm, margin_r=24 * mm)
    letterhead(d, ("Brennan Lowe Solicitors", ["For the Defendants", "Ref: BL/MED/7714"]))
    d.text(d.ml, "Part 35 questions to Dr A Expert (GP expert for the Claimant)", "Serif-B", 11, lead=18)
    d.text(d.ml, f"Date: {uk_date(date(2026, 1, 12))}", "Serif", 10.5, lead=18)
    for i, q in enumerate(qs, 1):
        d.para(d.ml, f"{i}. {q}", W - d.ml - d.mr, "Serif", 10.5, lead=14)
        d.y -= 6
    json.dump([{"number": i, "text_verbatim": q, "expected_kind": "goes_beyond" if i == 6 else "clarifies"}
               for i, q in enumerate(qs, 1)], open(os.path.join(EXP, "written_questions.json"), "w"), indent=1)
    return d.save(out)


def other_expert_report(out, pages):
    """12 claims with page citations, designed against records.pdf."""
    by_kind = {}
    for r in pages:
        for f in r["facts"]:
            by_kind.setdefault((f["kind"], f.get("planted")), []).append((r["page"], f))
    p7 = next(p for (p, f) in by_kind[("safety_netting", "7")])
    p10 = next(p for (p, f) in by_kind[("safety_netting", "10")])
    p_ferr = 402
    p_haem = next(p for (p, f) in by_kind[("letter", "8")])
    p_ooh = next(p for (p, f) in by_kind[("consultation", "13")])
    p_appt = next(p for (p, f) in by_kind[("letter", "12")])
    p_2ww = next(r["page"] for r in pages if r["doc_id"] == "gpc-2ww-2024")
    p_feb = next(r["page"] for r in pages for f in r["facts"] if f["kind"] == "consultation" and f.get("event_date") == "2024-02-12")
    p_jun = next(r["page"] for r in pages for f in r["facts"] if f["kind"] == "referral" and f.get("event_date") == "2024-06-24")
    claims = [
        ("The plan recorded on 15 May 2024 was 'review 3/52 if no better'", p7, "review 3/52 if no better", "matches"),
        ("A haematology Advice and Guidance response dated 10 May 2024 recommended FIT and GI investigation", p_haem,
         "Advice and Guidance response", "matches"),
        ("The out-of-hours clinician advised review by her own GP", p_ooh, "Advised to see own GP", "matches"),
        ("A routine gastroenterology referral was made on 24 June 2024", p_jun, "Routine referral to Gastroenterology", "matches"),
        ("The hospital acknowledged the routine referral on 30 August 2024", p_appt, "Referral received - appointment offered", "matches"),
        ("A suspected cancer referral was made on 20 August 2024", p_2ww, "SUSPECTED LOWER GASTROINTESTINAL CANCER REFERRAL", "matches"),
        ("Bloods were arranged at the first consultation on 12 February 2024", p_feb, "Bloods: FBC, ferritin", "matches"),
        ("A colonoscopy report dated 2 September 2024 describes the tumour", 512, "colonoscopy", "not_on_page"),
        ("A gastroenterology clinic letter of October 2024 confirms the referral was triaged routinely", 603, "triaged as routine", "not_on_page"),
        ("The ferritin in February 2024 was 19 ug/L", p_ferr, "19", "wrong_value"),
        ("No safety-netting advice was documented after 15 May 2024", p10, "Advised to call back or 111", "contradicted"),
        ("At the May consultation the GP noted that the patient had lost weight", p7, "Thinks lost a bit of weight", "matches_after_review"),
    ]
    d = Doc("other-expert-report", "opposing_report", margin_l=24 * mm, margin_r=24 * mm)
    d.text(d.ml, "REPORT ON BREACH OF DUTY - GENERAL PRACTICE", "Sans-B", 13, lead=20)
    d.text(d.ml, "Prepared for the Defendants by Dr B Other MRCGP. Date: 30 January 2026", "Sans", 10, lead=20)
    d.text(d.ml, "5. Summary of the records", "Sans-B", 11, lead=16)
    for i, (txt, page, _, _) in enumerate(claims, 1):
        d.para(d.ml, f"5.{i} {txt} [p.{page}].", W - d.ml - d.mr, "Sans", 10.5, lead=14)
        d.y -= 5
    d.text(d.ml, "6. Opinion", "Sans-B", 11, lead=16)
    d.para(d.ml, "6.1 In my opinion iron replacement with planned review was a reasonable approach on 15 May 2024. A reasonably "
                 "competent GP would not necessarily have referred on the suspected cancer pathway before August 2024.",
           W - d.ml - d.mr, "Sans", 10.5, lead=14)
    json.dump([{"number": f"5.{i}", "text_verbatim": f"{t} [p.{p}].", "cited_page": p, "cited_value": v, "expected_status": s}
               for i, (t, p, v, s) in enumerate(claims, 1)], open(os.path.join(EXP, "opposing_claims.json"), "w"), indent=1)
    return d.save(out)


def agenda(out):
    items = [
        "Do the experts agree that the ferritin result of 9 ug/L in February 2024 required action?",
        "Do the experts agree on what was recorded as the plan on 15 May 2024?",
        "By what date would a reasonably competent GP have referred on the suspected cancer pathway?",
        "Would a referral before 20 August 2024 probably have led to diagnosis before September 2024?",
    ]
    d = Doc("agenda", "agenda", margin_l=24 * mm, margin_r=24 * mm)
    d.text(d.ml, "AGENDA FOR DISCUSSION BETWEEN THE GP EXPERTS", "Sans-B", 12, lead=20)
    d.text(d.ml, "Agreed by the parties' solicitors, 20 February 2026", "Sans", 10, lead=18)
    for i, t in enumerate(items, 1):
        d.para(d.ml, f"{i}. {t}", W - d.ml - d.mr, "Sans", 10.5, lead=14)
        d.y -= 6
    json.dump({"items": items, "changed_position_item": 4,
               "draft_position": "probably would have been diagnosed in June 2024",
               "served_position": "more likely than not the cancer would not have been diagnosed before September (outside my expertise)"},
              open(os.path.join(EXP, "agenda.json"), "w"), indent=1)
    return d.save(out)


def notes_for_trial(out):
    txt = ("Key pages: 15 May plan 'review 3/52 if no better'. Ferritin 9 on 14 Feb, filed 'recheck 3 months'.\n"
           "Haematology A&G dated 10 May, received 22 May, filed 28 May.\n"
           "Gastro referral 24 June routine, acknowledged 30 August.\n"
           "The practice had a written policy requiring all low ferritin results to be discussed with a partner within 48 hours.\n")
    p = os.path.join(out, "notes-for-trial.txt")
    open(p, "w").write(txt)
    json.dump({"unsourced_sentence": "The practice had a written policy requiring all low ferritin results to be discussed "
                                     "with a partner within 48 hours."}, open(os.path.join(EXP, "notes_check.json"), "w"), indent=1)
    return p


def template_docx(out):
    import docx
    from docx.shared import Pt
    doc = docx.Document()
    styles = doc.styles
    for name, size in [("AE Heading", 14), ("AE Subheading", 12)]:
        st = styles.add_style(name, 1)
        st.font.name = "Georgia"
        st.font.size = Pt(size)
        st.font.bold = True
    body = styles.add_style("AE Body", 1)
    body.font.name = "Georgia"
    body.font.size = Pt(11)
    doc.add_paragraph("Dr A Expert MBBS FRCGP - Medico-legal reports", style="AE Subheading")
    for h in ["Introduction", "My instructions", "Documents I have read", "Background", "What happened",
              "My opinion", "Range of reasonable opinion", "Summary of conclusions", "Declaration and statement of truth"]:
        doc.add_paragraph(h, style="AE Heading")
        doc.add_paragraph("[text]", style="AE Body")
    fixed = datetime(2026, 1, 5, 9, 0, 0)
    props = doc.core_properties
    props.author = props.last_modified_by = "Dr A Expert"
    props.created = props.modified = fixed
    props.revision = 1
    os.makedirs(out, exist_ok=True)
    p = os.path.join(out, "template.docx")
    buf = io.BytesIO()
    doc.save(buf)
    _stable_zip(buf.getvalue(), p, fixed)
    return p


def _stable_zip(data: bytes, path: str, when: datetime) -> None:
    """Rewrite a zip (docx) with fixed entry times and order, so the same input gives the same bytes."""
    src = zipfile.ZipFile(io.BytesIO(data))
    with zipfile.ZipFile(path, "w", zipfile.ZIP_DEFLATED) as dst:
        for name in sorted(src.namelist(), key=lambda n: (n != "[Content_Types].xml", n)):
            info = zipfile.ZipInfo(name, date_time=when.timetuple()[:6])
            info.compress_type = zipfile.ZIP_DEFLATED
            info.external_attr = 0o644 << 16
            dst.writestr(info, src.read(name))


def trial_bundle(out, records_pdf, pages):
    """Re-ordered, re-paginated court bundle with 3 cited pages removed and 5 re-scanned."""
    rng = random.Random(11)
    reader = PdfReader(records_pdf)
    n = len(reader.pages)
    planted = json.load(open(os.path.join(EXP, "planted.json")))
    removed = {214, planted["13"][0], planted["11"][0]}     # BP review page, out-of-hours report, Feb ferritin entry
    # group pages by source as the court bundle does: GP records, then hospital, then other
    order_key = {"gp_record": 0, "ooh": 1, "hospital_letter": 2, "discharge_summary": 2, "result": 3, "letters": 4,
                 "referral": 0, "admin": 5}
    src = {r["page"]: r["source_type"] for r in pages}
    keep = [p for p in range(1, n + 1) if p not in removed]
    keep.sort(key=lambda p: (order_key.get(src[p], 6), p))
    rescans = set(rng.sample([p for p in keep if p in (118, 402) or rng.random() < 0.05][:40], 5))
    rescans.add(118)
    rescans = set(list(rescans)[:5])
    w = PdfWriter()
    conc = []
    tmp = os.path.join(out, "_trial_tmp")
    os.makedirs(tmp, exist_ok=True)
    for new_no, old in enumerate(keep, 1):
        if old in rescans:
            one = PdfWriter()
            one.add_page(reader.pages[old - 1])
            src_path = os.path.join(tmp, f"p{old}.pdf")
            one.write(open(src_path, "wb"))
            img = degrade(pdf_pages_to_images(src_path, 110)[0], rng, "poor")
            images_to_pdf([img], os.path.join(tmp, f"r{old}.pdf"))
            page = PdfReader(os.path.join(tmp, f"r{old}.pdf")).pages[0]
        else:
            page = reader.pages[old - 1]
        stamp = os.path.join(tmp, "stamp.pdf")
        c = canvas.Canvas(stamp, pagesize=A4)
        c.setFont("Helvetica-Bold", 14)
        c.drawRightString(A4[0] - 12 * mm, 10 * mm, f"C/{new_no}")
        c.save()
        page.merge_page(PdfReader(stamp).pages[0])
        w.add_page(page)
        conc.append({"old_page": old, "trial_page": f"C/{new_no}", "rescanned": old in rescans})
    p = os.path.join(out, "trial-bundle.pdf")
    w.write(open(p, "wb"))
    json.dump({"concordance": conc, "removed_pages": sorted(removed)}, open(os.path.join(EXP, "concordance.json"), "w"), indent=1)
    import shutil
    shutil.rmtree(tmp)
    return p


if __name__ == "__main__":
    if sys.argv[1:2] == ["--setup-fixtures"]:
        # Slice C's committed test inputs: python evals/generator/extras.py --setup-fixtures OUT EVAL_ROOT
        dest, root = sys.argv[2], sys.argv[3]
        EXP = os.path.join(root, "expected")
        print(letter_of_instruction(dest, exp=EXP))
        print(template_docx(dest))
        sys.exit(0)
    root = sys.argv[1] if len(sys.argv) > 1 else "out"
    OUT = os.path.join(root, "bundles", "synthetic-mrs-a")
    EXP = os.path.join(root, "expected")
    pages = json.load(open(os.path.join(EXP, "pages.json")))
    for fn in (letter_of_instruction, witness_statement, written_questions, agenda, notes_for_trial, template_docx):
        print(fn(OUT))
    print(other_expert_report(OUT, pages))
    print(trial_bundle(OUT, os.path.join(OUT, "records.pdf"), pages))
