"""Behaviour-test inputs for Mrs A: wrong-patient pages, non-PDF files, a foreign-language page,
a rotated page, versioned instructions, a court order, a second round of questions, a revised agenda
and a re-issued court bundle. All invented. Run after build.py and extras.py.

Usage: python evals/generator/inputs.py OUT_DIR   (OUT_DIR as used by build.py)

Writes OUT_DIR/bundles/synthetic-mrs-a/inputs/* (gitignored) and OUT_DIR/expected/inputs.json.
"""
from __future__ import annotations

import io
import json
import math
import os
import random
import struct
import sys
import wave
import zipfile
from datetime import date
from email.message import EmailMessage

from pypdf import PdfReader, PdfWriter
from reportlab.lib.units import mm

sys.path.insert(0, os.path.dirname(__file__))

from common import Doc, W, H, uk_date, long_date, fmt_nhs, nhs_number_valid  # noqa: E402
from story import PATIENT, PRACTICE  # noqa: E402
from scan import pdf_pages_to_images, degrade, images_to_pdf  # noqa: E402
from extras import letterhead, FIRM  # noqa: E402

# A second invented person whose pages were misfiled into Mrs A's record.
OTHER = {"title": "Mrs", "first": "Joan", "last": "Fairley", "dob": date(1958, 11, 2),
         "nhs": "9994471023", "address": ["3 Wren Lane", "Northam", "NM2 5JD"]}


def _valid_nhs(seed: str) -> str:
    base = seed[:9]
    for last in "0123456789":
        if nhs_number_valid(base + last):
            return base + last
    raise ValueError("no valid check digit")


OTHER["nhs"] = _valid_nhs(OTHER["nhs"])
assert OTHER["nhs"] != PATIENT["nhs"]


def gp_header(person):
    def hdr(d: Doc, n: int):
        c = d.c
        c.setFont("Sans-B", 9)
        c.drawString(d.ml, H - 14 * mm, f"{person['last'].upper()}, {person['first']} ({person['title']})")
        c.setFont("Sans", 8)
        c.drawString(d.ml, H - 18 * mm, f"DOB {uk_date(person['dob'])}   NHS No {fmt_nhs(person['nhs'])}   {', '.join(person['address'])}")
        c.drawRightString(W - d.mr, H - 14 * mm, PRACTICE["name"])
        c.drawRightString(W - d.mr, H - 18 * mm, "Consultations")
        return H - 30 * mm
    return hdr


def misfiled_pages(tmp):
    """4 pages of another patient's GP record, in the same print style."""
    d = Doc("misfiled", "primary_care_record", header=gp_header(OTHER))
    entries = [
        (date(2024, 3, 12), "Dr S Okafor", "Problem: Hypertension (review). BP 148/92. Amlodipine 5mg od continued. Rv 3/12."),
        (date(2024, 3, 12), "Practice nurse", "Bloods taken: U&E, HbA1c, lipids."),
        (date(2024, 4, 19), "Dr R Patel", "Results: HbA1c 51 mmol/mol. Discussed diet. Repeat 6/12."),
        (date(2024, 4, 19), "Reception", "Rx request: amlodipine 5mg. Issued."),
    ]
    for i, (dt, who, txt) in enumerate(entries):
        if i:
            d.page_break()
        d.text(d.ml, f"{dt.strftime('%d-%b-%Y')}  {who}", "Sans-B", 9.5, lead=14)
        d.para(d.ml, txt, W - d.ml - d.mr, "Sans", 9.5, lead=13)
    return d.save(tmp)


def mixed_records(out, records_pdf, tmp):
    """Pages 1-8 of Mrs A, then 4 misfiled pages, then pages 9-16: a realistic mid-bundle misfile."""
    mrs_a = PdfReader(records_pdf)
    other = PdfReader(misfiled_pages(tmp))
    w = PdfWriter()
    for i in range(8):
        w.add_page(mrs_a.pages[i])
    for p in other.pages:
        w.add_page(p)
    for i in range(8, 16):
        w.add_page(mrs_a.pages[i])
    path = os.path.join(out, "mixed-records.pdf")
    w.write(open(path, "wb"))
    return path


def welsh_letter(out):
    d = Doc("welsh-letter", "letters", margin_l=22 * mm, margin_r=22 * mm)
    d.text(d.ml, "Bwrdd Iechyd Dyffryn Northam (dychmygol)", "Serif-B", 12, lead=18)
    d.text(d.ml, long_date(date(2024, 10, 2)), "Serif", 10.5, lead=18)
    for p in [
        f"Annwyl Dr Patel,",
        f"Ail: {PATIENT['title']} {PATIENT['first']} {PATIENT['last']}, dyddiad geni {uk_date(PATIENT['dob'])}",
        "Diolch am gyfeirio'r claf i'r clinig. Fe'i gwelwyd heddiw. Mae'r driniaeth yn mynd yn ei blaen fel y cynlluniwyd "
        "ac nid oes angen newid ei meddyginiaeth ar hyn o bryd.",
        "Byddwn yn ei gweld eto ymhen chwe wythnos.",
        "Yn gywir,",
        "Dr E Rhys, Meddyg Ymgynghorol",
    ]:
        d.para(d.ml, p, W - d.ml - d.mr, "Serif", 10.5, lead=15)
        d.y -= 6
    return d.save(out)


def rotated_scan(out, records_pdf):
    r = PdfReader(records_pdf)
    w = PdfWriter()
    page = r.pages[117]          # p.118, the 15 May consultation
    page.rotate(90)
    w.add_page(page)
    path = os.path.join(out, "rotated-page.pdf")
    w.write(open(path, "wb"))
    return path


def photographed_page(out, records_pdf, tmp, rng):
    one = PdfWriter()
    one.add_page(PdfReader(records_pdf).pages[213])     # p.214, the February blood result
    src = os.path.join(tmp, "p214.pdf")
    one.write(open(src, "wb"))
    img = degrade(pdf_pages_to_images(src, 120)[0], rng, "poor").rotate(3.5, expand=True, fillcolor=210)
    path = os.path.join(out, "photographed-page.jpg")
    img.convert("RGB").save(path, quality=72)
    return path


def voicemail(out):
    """Five seconds of tones standing in for a recorded phone message (an exhibit, never read)."""
    path = os.path.join(out, "voicemail.wav")
    with wave.open(path, "wb") as f:
        f.setnchannels(1)
        f.setsampwidth(2)
        f.setframerate(8000)
        frames = b"".join(struct.pack("<h", int(6000 * math.sin(2 * math.pi * (440 if (i // 4000) % 2 else 523) * i / 8000)))
                          for i in range(8000 * 5))
        f.writeframes(frames)
    return path


def word_letter(out):
    from docx import Document
    doc = Document()
    doc.add_paragraph(PRACTICE["name"])
    doc.add_paragraph(long_date(date(2024, 8, 21)))
    doc.add_paragraph(f"Re: {PATIENT['title']} {PATIENT['first']} {PATIENT['last']}, DOB {uk_date(PATIENT['dob'])}")
    doc.add_paragraph("Dear Colleague, thank you for seeing this lady urgently with change in bowel habit, rectal bleeding "
                      "and iron deficiency anaemia (ferritin 9 ug/L in February 2024). Yours sincerely, Dr R Patel")
    path = os.path.join(out, "referral-letter.docx")
    doc.save(path)
    return path


def email_with_attachment(out, attachment):
    m = EmailMessage()
    m["From"] = "r.tennant@harlandprice.example"
    m["To"] = "expert@example.co.uk"
    m["Subject"] = "Ashworth: further records"
    m["Date"] = "Mon, 01 Dec 2025 10:14:00 +0000"
    m.set_content("Dear Dr Expert,\n\nPlease find attached the referral letter, which was missing from the bundle.\n\nRegards,\nR. Tennant")
    with open(attachment, "rb") as f:
        m.add_attachment(f.read(), maintype="application",
                         subtype="vnd.openxmlformats-officedocument.wordprocessingml.document",
                         filename=os.path.basename(attachment))
    path = os.path.join(out, "further-records.eml")
    open(path, "wb").write(bytes(m))
    return path


def records_zip(out, files):
    path = os.path.join(out, "records-batch.zip")
    with zipfile.ZipFile(path, "w") as z:
        for f in files:
            z.write(f, os.path.basename(f))
    return path


SUPP_Q = "If Mrs Ashworth had been examined on 15 May 2024, what findings would a reasonably competent general practitioner have expected to record?"


def supplementary_instructions(out):
    d = Doc("supplementary-letter-of-instruction", "loi", margin_l=24 * mm, margin_r=24 * mm)
    letterhead(d)
    d.text(d.ml, long_date(date(2025, 11, 20)), "Serif", 10.5, lead=18)
    d.text(d.ml, "Dear Dr Expert", "Serif", 10.5, lead=16)
    d.text(d.ml, "Margaret Ashworth v Brookvale Surgery: supplementary instructions", "Serif-B", 10.5, lead=18)
    for p in ["1. Further to our letter of 9 October 2025, we should be grateful if you would also answer the following question:",
              f"(e) {SUPP_Q}",
              "2. All other instructions are unchanged."]:
        d.para(d.ml, p, W - d.ml - d.mr, "Serif", 10.5, lead=14)
        d.y -= 6
    d.text(d.ml, "Yours sincerely", "Serif", 10.5, lead=26)
    d.text(d.ml, "R. Tennant, Partner, Harland & Price LLP", "Serif", 10.5)
    return d.save(out)


ORDER_TEXT = ("1. The experts in general practice shall discuss by 27 March 2026 and file a joint statement by 10 April 2026. "
              "2. The trial listed for 4 May 2026 is adjourned to the first available date after 1 October 2026, time estimate 3 days. "
              "3. Costs in the case.")


def court_order(out):
    d = Doc("court-order", "court_order", margin_l=24 * mm, margin_r=24 * mm)
    d.text(W / 2 - 40 * mm, "IN THE COUNTY COURT AT NORTHAM (invented)", "Serif-B", 11, lead=16)
    d.text(d.ml, "Claim No. N0RT-2291 (invented)", "Serif", 10, lead=16)
    d.text(d.ml, "MARGARET ASHWORTH  Claimant  -v-  BROOKVALE SURGERY  Defendant", "Serif-B", 10, lead=20)
    d.text(d.ml, f"ORDER made on {long_date(date(2026, 3, 10))} by District Judge Hale (invented)", "Serif", 10, lead=18)
    d.para(d.ml, ORDER_TEXT, W - d.ml - d.mr, "Serif", 10.5, lead=15)
    return d.save(out)


ROUND2 = ["Please confirm whether your answer to question 3 of our first set of questions relied on the practice task list.",
          "Please state the page reference for the entry recording that iron was started."]


def written_questions_round2(out):
    d = Doc("written-questions-round-2", "written_questions", margin_l=24 * mm, margin_r=24 * mm)
    letterhead(d, ("Brennan Lowe Solicitors", ["For the Defendants", "Ref: BL/MED/7714"]))
    d.text(d.ml, "Further written questions to Dr A Expert (second set)", "Serif-B", 11, lead=18)
    d.text(d.ml, f"Date: {uk_date(date(2026, 2, 16))}", "Serif", 10.5, lead=18)
    for i, q in enumerate(ROUND2, 1):
        d.para(d.ml, f"{i}. {q}", W - d.ml - d.mr, "Serif", 10.5, lead=14)
        d.y -= 6
    return d.save(out)


AGENDA_V2_ITEM3 = "By what date, if at all, would a reasonably competent GP have referred on the suspected cancer pathway?"


def agenda_v2(out):
    from extras import agenda as _agenda_items_source  # noqa: F401  (v1 built by extras.py)
    items = [
        "Do the experts agree that the ferritin result of 9 ug/L in February 2024 required action?",
        "Do the experts agree on what was recorded as the plan on 15 May 2024?",
        AGENDA_V2_ITEM3,
        "Would a referral before 20 August 2024 probably have led to diagnosis before September 2024?",
    ]
    d = Doc("agenda-v2", "agenda", margin_l=24 * mm, margin_r=24 * mm)
    d.text(d.ml, "AGENDA FOR DISCUSSION BETWEEN THE GP EXPERTS (REVISED)", "Sans-B", 12, lead=20)
    d.text(d.ml, "Revised by agreement, 24 March 2026 (item 3 reworded)", "Sans", 10, lead=18)
    for i, t in enumerate(items, 1):
        d.para(d.ml, f"{i}. {t}", W - d.ml - d.mr, "Sans", 10.5, lead=14)
        d.y -= 6
    return d.save(out)


def trial_bundle_v2(out, bundle_dir):
    """Re-issued court bundle: version 1 plus a 2-page supplementary section."""
    v1 = os.path.join(bundle_dir, "trial-bundle.pdf")
    if not os.path.exists(v1):
        print("skip trial-bundle-v2.pdf (run extras.py first)")
        return None
    r = PdfReader(v1)
    w = PdfWriter()
    for p in r.pages:
        w.add_page(p)
    d = Doc("supp", "court_bundle", margin_l=24 * mm, margin_r=24 * mm)
    for i, t in enumerate(["SUPPLEMENTARY BUNDLE: order of 10 March 2026", "SUPPLEMENTARY BUNDLE: revised agenda"], 1):
        if i > 1:
            d.page_break()
        d.text(d.ml, t, "Sans-B", 12, lead=18)
        d.text(W - d.mr - 20 * mm, f"S/{i}", "Sans-B", 14, y=10 * mm, advance=False)
    tmp = d.save(os.path.join(out, "_tmp"))
    for p in PdfReader(tmp).pages:
        w.add_page(p)
    path = os.path.join(out, "trial-bundle-v2.pdf")
    w.write(open(path, "wb"))
    return path


if __name__ == "__main__":
    root = sys.argv[1] if len(sys.argv) > 1 else "out"
    BUNDLE = os.path.join(root, "bundles", "synthetic-mrs-a")
    OUT = os.path.join(BUNDLE, "inputs")
    EXP = os.path.join(root, "expected")
    TMP = os.path.join(OUT, "_tmp")
    os.makedirs(TMP, exist_ok=True)
    rng = random.Random(23)
    records = os.path.join(BUNDLE, "records.pdf")
    made = {
        "mixed_records": mixed_records(OUT, records, TMP),
        "welsh_letter": welsh_letter(OUT),
        "rotated_page": rotated_scan(OUT, records),
        "photographed_page": photographed_page(OUT, records, TMP, rng),
        "voicemail": voicemail(OUT),
        "word_letter": word_letter(OUT),
    }
    made["email"] = email_with_attachment(OUT, made["word_letter"])
    made["zip"] = records_zip(OUT, [made["word_letter"], made["photographed_page"]])
    made["supplementary_instructions"] = supplementary_instructions(OUT)
    made["court_order"] = court_order(OUT)
    made["written_questions_round_2"] = written_questions_round2(OUT)
    made["agenda_v2"] = agenda_v2(OUT)
    made["trial_bundle_v2"] = trial_bundle_v2(OUT, BUNDLE)
    import shutil
    shutil.rmtree(TMP, ignore_errors=True)
    for k, v in made.items():
        print(k, v)
    expected = {
        "note": "Behaviour-test inputs. Paths are relative to bundles/synthetic-mrs-a/inputs/. Each case here is a separate small test case.",
        "wrong_patient": {
            "file": "mixed-records.pdf", "pages": 20, "misfiled_bundle_pages": [9, 10, 11, 12],
            "other_person": {"name": f"{OTHER['first']} {OTHER['last']}", "dob": OTHER["dob"].isoformat(), "nhs": OTHER["nhs"]},
            "expect": "pages 9-12 suggested for quarantine (identifiers don't match the case subject); never auto-hidden; "
                      "after confirmation numbers kept, placeholder shown, facts rejected, excluded from AI/search/citations/exports",
            "slice": "A", "scenario": "S5", "failure_mode": "F25"},
        "foreign_language": {"file": "welsh-letter.pdf", "language": "cy",
                             "expect": "flagged as not the pack language; no machine translation; certified-translation request offered",
                             "slice": "A", "failure_mode": "F37"},
        "rotation": {"file": "rotated-page.pdf", "same_content_as_page": 118,
                     "expect": "detected or rotated by the expert; OCR re-run; page number kept; verified facts re-mapped", "slice": "B"},
        "photographed_page": {"file": "photographed-page.jpg", "same_content_as_page": 214,
                              "expect": "converted to a PDF page; low confidence; offered as a clearer/alternative copy, never auto-merged",
                              "slice": "A", "scenario": "S6"},
        "exhibit": {"file": "voicemail.wav", "expect": "stored as an exhibit, listed as not readable, no page number, no AI call",
                    "slice": "A", "failure_mode": "F36"},
        "word_file": {"file": "referral-letter.docx", "expect": "converted to 1 page; text from the file (born digital)", "slice": "A"},
        "email": {"file": "further-records.eml", "expect": "email rendered as 1 page; the .docx attachment becomes its own document",
                  "slice": "A"},
        "zip": {"file": "records-batch.zip", "contains": ["referral-letter.docx", "photographed-page.jpg"],
                "expect": "unpacked with a manifest; each file handled as above", "slice": "A"},
        "instruction_version_2": {"file": "supplementary-letter-of-instruction.pdf", "new_question": SUPP_Q, "para_ref": "1(e)",
                                  "expect": "new instruction version; locked question set reopens with a warning; earlier links kept",
                                  "slice": "C", "scenario": "S7"},
        "court_order": {"file": "court-order.pdf", "made_on": "2026-03-10",
                        "sets": [{"kind": "expert_discussion", "due_on": "2026-03-27"},
                                 {"kind": "expert_meeting_document", "due_on": "2026-04-10"}],
                        "hearing": {"from": "2026-05-04", "status": "adjourned", "to_after": "2026-10-01", "days": 3},
                        "expect": "dates suggested from the order, confirmed by the expert; deadline_changes rows; trial pack and offline package marked stale",
                        "slice": "G", "scenario": "S10"},
        "written_questions_round_2": {"file": "written-questions-round-2.pdf", "questions": ROUND2,
                                      "expect": "second round with its own deadline and answers document", "slice": "G", "scenario": "S11"},
        "agenda_v2": {"file": "agenda-v2.pdf", "changed_items": [3], "new_item_3": AGENDA_V2_ITEM3,
                      "expect": "new agenda version; differences shown; out-for-signature returns to draft", "slice": "I"},
        "court_bundle_v2": {"file": "trial-bundle-v2.pdf", "added_section": "S", "added_pages": 2,
                            "expect": "supersedes version 1; new concordance; trial pack uses the latest confirmed version and old pack goes stale",
                            "slice": "I", "scenario": "S12", "failure_mode": "F35"},
    }
    json.dump(expected, open(os.path.join(EXP, "inputs.json"), "w"), indent=1)
    print("wrote", os.path.join(EXP, "inputs.json"))
